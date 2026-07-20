#!/usr/bin/env python3
"""Reproducibly build the canonical Bible 2.0 release (MD/DOCX/PDF/EPUB)."""
from __future__ import annotations
import hashlib, html, os, re, shutil, tempfile, zipfile
from datetime import datetime, timezone
from pathlib import Path

os.environ.setdefault("SOURCE_DATE_EPOCH", "1784505600")

import markdown
from bs4 import BeautifulSoup, NavigableString, Tag
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Mm, Pt, RGBColor
from weasyprint import HTML
import pikepdf
import fitz
from PIL import Image, ImageDraw

ROOT = Path(__file__).resolve().parent
MANUSCRIPT = ROOT / "manuscript"
OUTPUT = ROOT / "output"
FILES = [MANUSCRIPT / f for f in ["00_PROLOGUE.md", "01_BODY.md", "02_NAME.md", "03_FIRE.md", "04_FORGE.md", "05_COVENANT.md", "06_CROWN.md", "07_VEIL.md", "08_HARVEST.md", "09_GATES.md", "10_SEAL.md"]]
STEM = "BIBLE_2_0_0_TO_1_RELEASE"
EPOCH = (2026, 7, 20, 0, 0, 0)
TITLE = "БИБЛИЯ 2.0"
SUBTITLE = "0→1 · Кодекс перехода"
AUTHOR = "Каноническая редакция"
LANG = "ru-RU"

CSS = r"""
@page { size: A5; margin: 17mm 15mm 19mm 18mm; @top-left { content: string(book); font: 7.5pt 'DejaVu Sans'; color:#6d6964; border-bottom:.3pt solid #b9b1a3 } @top-right { content:'0→1'; font:7.5pt 'DejaVu Sans'; color:#6d6964; border-bottom:.3pt solid #b9b1a3 } @bottom-outer { content: counter(page); font:8pt 'DejaVu Sans'; color:#6d6964 } }
@page:left { margin-left:15mm; margin-right:18mm } @page:right { margin-left:18mm; margin-right:15mm }
@page cover { margin:0; @top-left{content:none} @top-right{content:none} @bottom-outer{content:none} }
@page opening { @top-left{content:none} @top-right{content:none} }
html { font-family:'DejaVu Serif',serif; font-size:9.6pt; line-height:1.32; color:#171b22; }
body { margin:0; hyphens:auto; }
.cover { page:cover; height:209.5mm; margin:0; padding:42mm 18mm; box-sizing:border-box; background:#171b22; color:#f4f0e6; text-align:center; break-after:page; }
.cover .sigil {font:42pt 'DejaVu Sans'; color:#b6643a; margin:12mm 0}.cover h1{page:auto;break-before:auto;font:27pt 'DejaVu Sans'; letter-spacing:2pt;margin:0}.cover h2{page:auto;break-before:auto;font:italic 14pt 'DejaVu Serif';font-weight:normal;border:0}.cover p{margin-top:48mm;font:8pt 'DejaVu Sans';color:#b9b1a3}
.front-title { text-align:center; break-after:page; padding-top:45mm }.front-title h1{font:25pt 'DejaVu Sans'}.front-title .formula{font:italic 13pt 'DejaVu Serif'; color:#b6643a;margin-top:12mm}
.toc { break-after:page }.toc h1,.back h1{font:22pt 'DejaVu Sans'}.toc ul{list-style:none;padding:0}.toc li{margin:2.5mm 0}.toc a{color:#171b22;text-decoration:none}.toc a::after{content:leader('.') target-counter(attr(href),page);float:right}
h1 { string-set:book content(); page:opening; break-before:page; font:22pt/1.12 'DejaVu Sans'; color:#244b78; margin:0 0 12mm; padding-top:16mm; }
h2 { break-before:page; font:15.5pt/1.18 'DejaVu Sans'; color:#171b22; margin:0 0 7mm; padding-top:5mm; break-after:avoid; }
h3 { font:11.5pt 'DejaVu Sans'; color:#244b78; margin:6mm 0 2mm; break-after:avoid; }
p { margin:0 0 2.2mm; text-align:justify; orphans:3; widows:3; } blockquote{margin:4mm 7mm;font-style:italic;color:#34383d} blockquote p{text-align:left}
pre,code{font:8.3pt 'DejaVu Sans Mono';white-space:pre-wrap} hr{border:0;border-top:.4pt solid #b9b1a3;margin:7mm 20mm} ul,ol{padding-left:7mm} li{margin-bottom:1.2mm} strong{font-weight:bold}.back{break-before:page}.status{border-left:2pt solid #b6643a;padding-left:4mm}.colophon{text-align:center;color:#6d6964;font-size:8pt;margin-top:20mm}
"""

BACK_MD = """# ПОСЛЕСЛОВИЕ: КАРТА СТАТУСОВ ИСТОЧНИКОВ

Эта книга различает четыре режима речи, не разрывая ими повествование.

- **Факт** — проверяемое историческое, естественно-научное или текстологическое утверждение.
- **Спорная версия** — атрибуция или толкование, о котором источники не дают единого ответа.
- **Авторская модель** — рабочая конструкция Кодекса: Ось, Круг, Спираль и их соответствия.
- **Миф** — образ, сохранённый ради сравнительного смысла, а не выданный за буквальную хронику.

Символические сближения в книге — сравнительное чтение, а не заявление о едином тайном происхождении традиций. Этимологии относятся к языковой истории; созвучия, гематрические и нумерологические мосты — к авторской поэтике, если прямо не указано иное. Полная рабочая карта источников остаётся за пределами основного текста.

---

*Каноническая редакция · 20 июля 2026 · UTF-8 · A5*
"""

def slug(s: str, used: set[str]) -> str:
    base = re.sub(r"[^\wа-яё]+", "-", s.lower(), flags=re.I).strip("-") or "section"
    out=base; n=2
    while out in used: out=f"{base}-{n}"; n+=1
    used.add(out); return out

def canonical_markdown() -> tuple[str,list[tuple[int,str,str]]]:
    sources=[p.read_text(encoding="utf-8").rstrip()+"\n" for p in FILES]
    headings=[]; used=set()
    for src in sources:
        for marks,title in re.findall(r"^(#{1,2})\s+(.+)$",src,re.M):
            headings.append((len(marks),title.strip(),slug(title.strip(),used)))
    toc=["# СОДЕРЖАНИЕ",""]
    for level,title,anchor in headings:
        toc.append(f"{'  ' if level==2 else ''}- [{title}](#{anchor})")
    front=f"# {TITLE}\n\n## {SUBTITLE}\n\n{AUTHOR}\n\nРедакция 20 июля 2026\n"
    body="\n\n<div style=\"page-break-after: always\"></div>\n\n".join(s.rstrip() for s in sources)
    return front+"\n\n"+"\n".join(toc)+"\n\n"+body+"\n\n"+BACK_MD, headings

def body_html(md_text:str, headings):
    # Only canonical manuscript is rendered here; front matter is custom.
    body="\n\n".join(p.read_text(encoding='utf-8').rstrip() for p in FILES)+"\n\n"+BACK_MD
    raw=markdown.markdown(body,extensions=['extra','sane_lists'])
    soup=BeautifulSoup(raw,'html.parser'); used=set()
    for h in soup.find_all(['h1','h2']): h['id']=slug(h.get_text(' ',strip=True),used)
    toc=''.join(f'<li class="l{lev}"><a href="#{anchor}">{html.escape(title)}</a></li>' for lev,title,anchor in headings)
    return f'''<!doctype html><html lang="ru"><head><meta charset="utf-8"><title>{TITLE}: {SUBTITLE}</title><meta name="author" content="{AUTHOR}"><style>{CSS}</style></head><body>
<section class="cover"><h1>{TITLE}</h1><div class="sigil">○│•</div><h2>{SUBTITLE}</h2><p>КАНОНИЧЕСКАЯ РЕДАКЦИЯ · 2026</p></section>
<section class="front-title"><h1>{TITLE}</h1><div class="formula">Ноль вмещает · Единица отвечает</div><p>{SUBTITLE}</p></section>
<nav class="toc"><h1>Содержание</h1><ul>{toc}</ul></nav>{soup}</body></html>'''

def set_cell_font(run, name="DejaVu Serif"):
    run.font.name=name; run._element.rPr.rFonts.set(qn('w:eastAsia'),name); run._element.rPr.rFonts.set(qn('w:cs'),name)

def add_field(paragraph, code):
    run=paragraph.add_run(); begin=OxmlElement('w:fldChar'); begin.set(qn('w:fldCharType'),'begin'); instr=OxmlElement('w:instrText'); instr.set(qn('xml:space'),'preserve'); instr.text=code; sep=OxmlElement('w:fldChar'); sep.set(qn('w:fldCharType'),'separate'); end=OxmlElement('w:fldChar'); end.set(qn('w:fldCharType'),'end'); run._r.extend([begin,instr,sep,end])

def build_docx(html_text, target):
    d=Document(); sec=d.sections[0]; sec.page_height=Mm(210);sec.page_width=Mm(148);sec.top_margin=Mm(17);sec.bottom_margin=Mm(19);sec.left_margin=Mm(18);sec.right_margin=Mm(15);sec.header_distance=Mm(7);sec.footer_distance=Mm(8);sec.different_first_page_header_footer=True
    styles=d.styles
    normal=styles['Normal']; normal.font.name='DejaVu Serif';normal.font.size=Pt(10);normal.paragraph_format.space_after=Pt(4);normal.paragraph_format.line_spacing=1.16
    for sn,size,color in [('Title',29,'171B22'),('Heading 1',22,'244B78'),('Heading 2',16,'171B22'),('Heading 3',11.5,'244B78')]:
        st=styles[sn];st.font.name='DejaVu Sans';st.font.size=Pt(size);st.font.color.rgb=RGBColor.from_string(color);st.font.bold=True
    for st in [styles['Heading 1'],styles['Heading 2']]: st.paragraph_format.page_break_before=True;st.paragraph_format.keep_with_next=True
    styles['Heading 3'].paragraph_format.keep_with_next=True
    # Header/footer furniture.
    hp=sec.header.paragraphs[0]; hp.text=f"{TITLE}   ·   0→1";hp.alignment=WD_ALIGN_PARAGRAPH.CENTER
    fp=sec.footer.paragraphs[0];fp.alignment=WD_ALIGN_PARAGRAPH.CENTER;add_field(fp,'PAGE')
    d.core_properties.title=TITLE;d.core_properties.subject=SUBTITLE;d.core_properties.author=AUTHOR;d.core_properties.language=LANG;d.core_properties.keywords='0→1, мифопоэтика, символизм, современная светская Библия';fixed=datetime(2026,7,20,tzinfo=timezone.utc);d.core_properties.created=fixed;d.core_properties.modified=fixed
    p=d.add_paragraph();p.style='Title';p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.add_run(TITLE);d.add_paragraph('○│•').alignment=WD_ALIGN_PARAGRAPH.CENTER; p=d.add_paragraph(SUBTITLE);p.alignment=WD_ALIGN_PARAGRAPH.CENTER; d.add_page_break()
    d.add_heading('Содержание',0); p=d.add_paragraph(); add_field(p,'TOC \\o "1-2" \\h \\z \\u'); d.add_page_break()
    soup=BeautifulSoup(html_text,'html.parser')
    content=soup.body
    # Skip custom cover/front/toc already constructed.
    for node in list(content.children)[3:]:
        if not isinstance(node,Tag): continue
        if node.name in ('h1','h2','h3'):
            p=d.add_paragraph(node.get_text(' ',strip=True),style={'h1':'Heading 1','h2':'Heading 2','h3':'Heading 3'}[node.name])
        elif node.name=='p':
            p=d.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
            for child in node.children:
                if isinstance(child,NavigableString): p.add_run(str(child))
                elif isinstance(child,Tag):
                    r=p.add_run(child.get_text());r.bold=child.name in ('strong','b');r.italic=child.name in ('em','i')
        elif node.name=='blockquote':
            p=d.add_paragraph(node.get_text(' ',strip=True));p.style='Quote'
        elif node.name in ('ul','ol'):
            for li in node.find_all('li',recursive=False): d.add_paragraph(li.get_text(' ',strip=True),style='List Bullet' if node.name=='ul' else 'List Number')
        elif node.name=='hr': d.add_paragraph('•  •  •').alignment=WD_ALIGN_PARAGRAPH.CENTER
        elif node.name=='pre': d.add_paragraph(node.get_text(),style='No Spacing')
    for p in d.paragraphs:
        for run in p.runs: set_cell_font(run,'DejaVu Sans' if p.style.name.startswith(('Heading','Title')) else 'DejaVu Serif')
    d.save(target); normalize_zip(target)

def normalize_zip(path):
    with zipfile.ZipFile(path,'r') as z: items=[(i.filename,z.read(i.filename),i.compress_type) for i in z.infolist()]
    tmp=path.with_suffix(path.suffix+'.tmp')
    with zipfile.ZipFile(tmp,'w') as z:
        for name,data,ctype in sorted(items):
            zi=zipfile.ZipInfo(name,EPOCH);zi.compress_type=ctype;zi.external_attr=0o644<<16;z.writestr(zi,data)
    tmp.replace(path)

def build_epub(html_text,target):
    soup=BeautifulSoup(html_text,'html.parser'); chapters=[]
    nodes=[]; title='Вступление'
    for n in soup.body.children:
        if not isinstance(n,Tag) or 'cover' in n.get('class',[]) or 'front-title' in n.get('class',[]) or n.name=='nav': continue
        if n.name=='h1' and nodes:
            chapters.append((title,nodes));nodes=[]
        if n.name=='h1':title=n.get_text(' ',strip=True)
        nodes.append(str(n))
    if nodes:chapters.append((title,nodes))
    uid='urn:sha256:'+hashlib.sha256(''.join(p.read_text(encoding='utf-8') for p in FILES).encode()).hexdigest()
    files={'mimetype':b'application/epub+zip','META-INF/container.xml':b'''<?xml version="1.0"?><container version="1.0" xmlns="urn:oasis:names:tc:opendocument:xmlns:container"><rootfiles><rootfile full-path="EPUB/package.opf" media-type="application/oebps-package+xml"/></rootfiles></container>'''}
    nav=[];manifest=[];spine=[]
    for i,(title,nodes) in enumerate(chapters,1):
        fn=f'ch{i:02}.xhtml'; nav.append(f'<li><a href="{fn}">{html.escape(title)}</a></li>');manifest.append(f'<item id="c{i}" href="{fn}" media-type="application/xhtml+xml"/>');spine.append(f'<itemref idref="c{i}"/>')
        x=f'''<?xml version="1.0" encoding="utf-8"?><html xmlns="http://www.w3.org/1999/xhtml" xml:lang="ru"><head><title>{html.escape(title)}</title><link rel="stylesheet" href="style.css" type="text/css"/></head><body>{''.join(nodes)}</body></html>'''
        files['EPUB/'+fn]=x.encode()
    files['EPUB/style.css']=b"body{font-family:serif;line-height:1.45;margin:5%}h1{page-break-before:always}h2{page-break-before:always;color:#244b78}p{text-align:justify}blockquote{font-style:italic}"
    navx=f'''<?xml version="1.0" encoding="utf-8"?><html xmlns="http://www.w3.org/1999/xhtml" xmlns:epub="http://www.idpf.org/2007/ops" xml:lang="ru"><head><title>Содержание</title></head><body><nav epub:type="toc" id="toc"><h1>Содержание</h1><ol>{''.join(nav)}</ol></nav></body></html>''';files['EPUB/nav.xhtml']=navx.encode()
    opf=f'''<?xml version="1.0" encoding="utf-8"?><package xmlns="http://www.idpf.org/2007/opf" unique-identifier="uid" version="3.0" xml:lang="ru"><metadata xmlns:dc="http://purl.org/dc/elements/1.1/"><dc:identifier id="uid">{uid}</dc:identifier><dc:title>{TITLE}: {SUBTITLE}</dc:title><dc:language>ru</dc:language><dc:creator>{AUTHOR}</dc:creator><meta property="dcterms:modified">2026-07-20T00:00:00Z</meta></metadata><manifest><item id="nav" href="nav.xhtml" media-type="application/xhtml+xml" properties="nav"/><item id="css" href="style.css" media-type="text/css"/>{''.join(manifest)}</manifest><spine>{''.join(spine)}</spine></package>''';files['EPUB/package.opf']=opf.encode()
    with zipfile.ZipFile(target,'w') as z:
        zi=zipfile.ZipInfo('mimetype',EPOCH);zi.compress_type=zipfile.ZIP_STORED;z.writestr(zi,files.pop('mimetype'))
        for name,data in sorted(files.items()):zi=zipfile.ZipInfo(name,EPOCH);zi.compress_type=zipfile.ZIP_DEFLATED;zi.external_attr=0o644<<16;z.writestr(zi,data)

def build_visual_qa(pdf_target: Path) -> None:
    """Render every PDF page and assemble deterministic 20-page contact sheets."""
    pages_dir = ROOT / "qa" / "pdf_pages"
    sheets_dir = ROOT / "qa" / "contact_sheets"
    shutil.rmtree(pages_dir, ignore_errors=True)
    shutil.rmtree(sheets_dir, ignore_errors=True)
    pages_dir.mkdir(parents=True)
    sheets_dir.mkdir(parents=True)
    pdf = fitz.open(pdf_target)
    rendered = []
    for number, page in enumerate(pdf, 1):
        target = pages_dir / f"p{number:04}.png"
        page.get_pixmap(matrix=fitz.Matrix(1.2, 1.2), alpha=False).save(target)
        rendered.append(target)
    for start in range(0, len(rendered), 20):
        sheet = Image.new("RGB", (1300, 1480), "#777777")
        draw = ImageDraw.Draw(sheet)
        for offset, target in enumerate(rendered[start:start + 20]):
            with Image.open(target) as source:
                source.thumbnail((240, 340))
                x = (offset % 5) * 260 + 10
                y = (offset // 5) * 370 + 20
                sheet.paste(source, (x, y))
            draw.text((x, y - 16), target.stem, fill="white")
        sheet.save(sheets_dir / f"sheet_{start // 20 + 1:02}.jpg", quality=90)
    print(f"visual QA\t{len(rendered)} pages\t{(len(rendered) + 19) // 20} sheets")

def main():
    OUTPUT.mkdir(exist_ok=True)
    md_text,headings=canonical_markdown(); (OUTPUT/f'{STEM}.md').write_text(md_text,encoding='utf-8',newline='\n')
    html_text=body_html(md_text,headings)
    pdf_target=OUTPUT/f'{STEM}.pdf'
    HTML(string=html_text,base_url=str(ROOT)).write_pdf(pdf_target)
    with pikepdf.open(pdf_target) as pdf:
        for key in ['/CreationDate','/ModDate']:
            if key in pdf.docinfo: del pdf.docinfo[key]
        pdf.save(pdf_target.with_suffix('.pdf.tmp'), deterministic_id=True, static_id=True, normalize_content=True)
    pdf_target.with_suffix('.pdf.tmp').replace(pdf_target)
    build_docx(html_text,OUTPUT/f'{STEM}.docx')
    build_epub(html_text,OUTPUT/f'{STEM}.epub')
    build_visual_qa(pdf_target)
    for p in sorted(OUTPUT.glob(STEM+'.*')): print(f'{p.name}\t{p.stat().st_size}\t{hashlib.sha256(p.read_bytes()).hexdigest()}')
if __name__=='__main__':main()
