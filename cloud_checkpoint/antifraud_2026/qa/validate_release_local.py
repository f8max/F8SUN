#!/usr/bin/env python3
"""Independent, read-only release gate for the anti-fraud guide.

Normal validation never modifies release files.  ``--self-test`` deliberately
creates disposable negative fixtures under the operating-system temp folder and
proves that the gate rejects them.
"""

from __future__ import annotations

import argparse
import hashlib
import json
import posixpath
import re
import shutil
import subprocess
import sys
import tempfile
import unicodedata
import zipfile
from dataclasses import dataclass, field
from pathlib import Path
from typing import Iterable
from urllib.parse import unquote, urlsplit
from xml.etree import ElementTree as ET


BASE_NAME = "ANTIFRAUD_ULTIMATE_GUIDE_2026_RU"
A5_POINTS = (148 * 72 / 25.4, 210 * 72 / 25.4)
OOXML_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
OOXML_REL_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
PKG_REL_NS = "http://schemas.openxmlformats.org/package/2006/relationships"
DC_NS = "http://purl.org/dc/elements/1.1/"
CONTAINER_NS = "urn:oasis:names:tc:opendocument:xmlns:container"

REQUIRED_HEADING_PREFIXES = (
    "Деньги",
    "Карта угроз",
    "Fraud journey",
    "Эталонная",
    "ML без магии",
    "Bot management",
    "Практические чек-листы",
    "Глоссарий",
)

RAW_MERMAID_DECLARATION_RE = re.compile(
    r"(?mix)^\s*(?:"
    r"(?:flowchart|graph)\s+(?:TB|TD|BT|RL|LR)\b|"
    r"(?:sequenceDiagram|classDiagram|stateDiagram(?:-v2)?|erDiagram|journey|gantt|"
    r"mindmap|timeline)\s*(?:%%.*)?$|"
    r"pie(?:\s+showData)?(?:\s+title\b.*)?\s*$"
    r")",
)
RAW_TEX_RE = re.compile(r"\\(?:cdot|mid|frac|sum|operatorname)\b|^\s*\\[\[\]]\s*$", re.M)

REQUIRED_MANIFEST_INPUT_ROLES = (
    "source",
    "source_ledger",
    "claim_map",
    "builder",
    "requirements",
    "package_json",
    "package_lock",
    "puppeteer_config",
)
REQUIRED_MANIFEST_TOOLS = ("python", "python_packages", "weasyprint", "mmdc")


@dataclass
class Reporter:
    quiet: bool = False
    passes: int = 0
    warnings: int = 0
    failures: int = 0
    records: list[tuple[str, str, str]] = field(default_factory=list)

    def _emit(self, level: str, scope: str, message: str) -> None:
        self.records.append((level, scope, message))
        if not self.quiet:
            print(f"{level:<4} [{scope}] {message}")

    def passed(self, scope: str, message: str) -> None:
        self.passes += 1
        self._emit("PASS", scope, message)

    def warn(self, scope: str, message: str) -> None:
        self.warnings += 1
        self._emit("WARN", scope, message)

    def fail(self, scope: str, message: str) -> None:
        self.failures += 1
        self._emit("FAIL", scope, message)

    def summary(self, strict_warnings: bool = False) -> int:
        status = "FAIL" if self.failures or (strict_warnings and self.warnings) else "PASS"
        print(
            f"{status}: {self.passes} passed, {self.warnings} warnings, "
            f"{self.failures} failures"
        )
        return 1 if status == "FAIL" else 0


@dataclass
class ReleaseContext:
    book_path: Path
    ledger_path: Path
    book_text: str = ""
    ledger_text: str = ""
    plain_text: str = ""
    metadata: dict[str, str] = field(default_factory=dict)
    headings: list[tuple[int, str]] = field(default_factory=list)
    citations: set[str] = field(default_factory=set)
    ledger_ids: set[str] = field(default_factory=set)
    ledger_urls: dict[str, str] = field(default_factory=dict)
    mermaid_blocks: list[str] = field(default_factory=list)
    expected_tables: int = 0
    cutoff_date: str = ""
    word_count: int = 0

    @property
    def title(self) -> str:
        return self.metadata.get("title", BASE_NAME)

    @property
    def author(self) -> str:
        return self.metadata.get("author", "")

    @property
    def h1_titles(self) -> list[str]:
        return [title for level, title in self.headings if level == 1]


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def local_name(tag: str) -> str:
    return tag.rsplit("}", 1)[-1]


def normalize_visible(value: str) -> str:
    """Normalize user-visible metadata without changing its semantic content."""
    return re.sub(r"\s+", " ", unicodedata.normalize("NFKC", value)).strip().casefold()


def normalize_heading(value: str) -> str:
    """Normalize a heading while ignoring a leading chapter/section number."""
    without_number = re.sub(r"^\s*(?:глава\s+)?\d+(?:\.\d+)*[.)]?\s*", "", value, flags=re.I)
    return normalize_visible(without_number)


def contains_raw_mermaid(text: str, markup: str = "") -> bool:
    """Detect an unrendered Mermaid declaration, not ordinary prose arrows/rules."""
    return (
        "language-mermaid" in markup.casefold()
        or RAW_MERMAID_DECLARATION_RE.search(text) is not None
    )


def element_text(element: ET.Element | None) -> str:
    if element is None:
        return ""
    return " ".join("".join(element.itertext()).split())


def read_utf8(path: Path, reporter: Reporter, scope: str) -> str | None:
    if not path.is_file():
        reporter.fail(scope, f"file not found: {path}")
        return None
    try:
        return path.read_text(encoding="utf-8-sig")
    except (OSError, UnicodeError) as exc:
        reporter.fail(scope, f"cannot read UTF-8 file {path}: {exc}")
        return None


def parse_frontmatter(text: str) -> tuple[dict[str, str], str]:
    metadata: dict[str, str] = {}
    normalized = text.replace("\r\n", "\n").replace("\r", "\n")
    if not normalized.startswith("---\n"):
        return metadata, normalized
    end = normalized.find("\n---\n", 4)
    if end < 0:
        return metadata, normalized
    for line in normalized[4:end].splitlines():
        if ":" not in line or line.lstrip().startswith("#"):
            continue
        key, value = line.split(":", 1)
        metadata[key.strip()] = value.strip().strip("\"'")
    return metadata, normalized[end + 5 :]


def analyze_markdown_body(
    body: str,
) -> tuple[list[tuple[int, str]], list[tuple[str, str]], int, bool, str]:
    headings: list[tuple[int, str]] = []
    fences: list[tuple[str, str]] = []
    outside_lines: list[str] = []
    fence_marker = ""
    fence_language = ""
    fence_lines: list[str] = []
    table_count = 0
    previous_outside = ""

    for line in body.splitlines():
        fence_match = re.match(r"^\s*(`{3,}|~{3,})\s*([^\s`]*)", line)
        if fence_marker:
            if fence_match and fence_match.group(1).startswith(fence_marker[0]):
                fences.append((fence_language.casefold(), "\n".join(fence_lines)))
                fence_marker = ""
                fence_language = ""
                fence_lines = []
            else:
                fence_lines.append(line)
            continue
        if fence_match:
            fence_marker = fence_match.group(1)
            fence_language = fence_match.group(2)
            fence_lines = []
            continue

        outside_lines.append(line)
        heading = re.match(r"^(#{1,6})\s+(.+?)\s*#*\s*$", line)
        if heading:
            headings.append((len(heading.group(1)), heading.group(2).strip()))

        if re.match(r"^\s*\|?(?:\s*:?-{3,}:?\s*\|)+\s*$", line) and previous_outside.lstrip().startswith("|"):
            table_count += 1
        previous_outside = line

    unclosed = bool(fence_marker)
    if unclosed:
        fences.append((fence_language.casefold(), "\n".join(fence_lines)))
    return headings, fences, table_count, unclosed, "\n".join(outside_lines)


def markdown_to_plain(text: str) -> str:
    _, body = parse_frontmatter(text)
    _, _, _, _, outside = analyze_markdown_body(body)
    plain = re.sub(r"<[^>]+>", " ", outside)
    plain = re.sub(r"!\[([^]]*)\]\([^)]*\)", r"\1", plain)
    plain = re.sub(r"\[([^]]+)\]\([^)]*\)", r"\1", plain)
    plain = re.sub(r"https?://[^\s)>]+", " ", plain, flags=re.I)
    plain = re.sub(r"(?m)^\s*#{1,6}\s+", "", plain)
    plain = re.sub(r"(?m)^\s*(?:[-*+] |\d+\. )", "", plain)
    plain = plain.replace("|", " ")
    plain = re.sub(r"[*_`~\\]", "", plain)
    plain = re.sub(r"\s+", " ", plain)
    return plain.strip()


def count_unicode_words(text: str) -> int:
    """Count letter-bearing Unicode word tokens, excluding pure numbers."""
    tokens = re.findall(r"(?u)(?<!\w)[^\W_]+(?:[-’'][^\W_]+)*(?!\w)", text)
    return sum(1 for token in tokens if any(character.isalpha() for character in token))


def extract_cutoff(text: str) -> str:
    match = re.search(r"\b(20\d{2}-\d{2}-\d{2})\b", text[:4000])
    return match.group(1) if match else ""


def validate_markdown_and_ledger(
    book: Path,
    ledger: Path,
    reporter: Reporter,
    min_book_words: int = 16_000,
) -> ReleaseContext:
    scope = "markdown"
    context = ReleaseContext(book_path=book, ledger_path=ledger)
    book_text = read_utf8(book, reporter, scope)
    ledger_text = read_utf8(ledger, reporter, "ledger")
    if book_text is None or ledger_text is None:
        return context

    context.book_text = book_text
    context.ledger_text = ledger_text
    metadata, body = parse_frontmatter(book_text)
    context.metadata = metadata
    context.plain_text = markdown_to_plain(book_text)
    context.word_count = count_unicode_words(context.plain_text)
    headings, fences, table_count, unclosed, outside = analyze_markdown_body(body)
    context.headings = headings
    context.expected_tables = table_count
    context.mermaid_blocks = [content for language, content in fences if language == "mermaid"]
    context.citations = set(re.findall(r"\[(S\d{2,})\]", outside))
    context.cutoff_date = extract_cutoff(book_text)

    if not metadata.get("title") or not metadata.get("author") or not metadata.get("date"):
        reporter.fail(scope, "front matter must contain non-empty title, author and date")
    else:
        reporter.passed(scope, "front matter has title, author and date")
    if unclosed:
        reporter.fail(scope, "unclosed fenced code block")
    else:
        reporter.passed(scope, f"all fenced blocks close; Mermaid blocks: {len(context.mermaid_blocks)}")
    if len(context.mermaid_blocks) < 3:
        reporter.fail(scope, "fewer than three Mermaid source diagrams")
    for number, block in enumerate(context.mermaid_blocks, 1):
        first = next((line.strip() for line in block.splitlines() if line.strip()), "")
        if not first or not re.match(
            r"^(?:flowchart|graph|sequenceDiagram|classDiagram|stateDiagram|erDiagram|journey|gantt|pie|mindmap|timeline)\b",
            first,
            re.I,
        ):
            reporter.fail(scope, f"Mermaid block {number} has no recognized diagram declaration")

    normalized_titles = [normalize_heading(title) for _, title in headings]
    missing_headings = [
        required
        for required in REQUIRED_HEADING_PREFIXES
        if not any(normalize_heading(required) in title for title in normalized_titles)
    ]
    if missing_headings:
        reporter.fail(scope, f"missing required headings: {', '.join(missing_headings)}")
    else:
        reporter.passed(scope, f"required structure present ({len(headings)} headings, {table_count} tables)")
    if not context.cutoff_date:
        reporter.fail(scope, "ISO cutoff date is absent near the beginning of the manuscript")
    else:
        reporter.passed(scope, f"cutoff date: {context.cutoff_date}")
    if len(context.plain_text) < 20_000:
        reporter.fail(scope, f"manuscript is unexpectedly short ({len(context.plain_text)} plain characters)")
    else:
        reporter.passed(scope, f"manuscript volume: {len(context.plain_text)} plain characters")
    if context.word_count < min_book_words:
        reporter.fail(
            scope,
            f"minimum volume gate: {context.word_count:,} Unicode words < required {min_book_words:,}",
        )
    else:
        reporter.passed(
            scope,
            f"minimum volume gate: {context.word_count:,} Unicode words >= required {min_book_words:,}",
        )

    ledger_ids: list[str] = []
    ledger_urls: dict[str, str] = {}
    for line in ledger_text.splitlines():
        match = re.match(r"^\|\s*(S\d{2,})\s*\|", line)
        if not match:
            continue
        source_id = match.group(1)
        urls = re.findall(r"https://[^\s|<>]+", line)
        ledger_ids.append(source_id)
        if len(urls) != 1:
            reporter.fail("ledger", f"{source_id} must have exactly one HTTPS URL")
        else:
            ledger_urls[source_id] = urls[0].rstrip(".,;:")

    context.ledger_ids = set(ledger_ids)
    context.ledger_urls = ledger_urls
    duplicate_ids = sorted({source_id for source_id in ledger_ids if ledger_ids.count(source_id) > 1})
    duplicate_urls = sorted({url for url in ledger_urls.values() if list(ledger_urls.values()).count(url) > 1})
    if duplicate_ids:
        reporter.fail("ledger", f"duplicate source IDs: {', '.join(duplicate_ids)}")
    if duplicate_urls:
        reporter.fail("ledger", f"duplicate source URLs: {', '.join(duplicate_urls)}")
    if len(context.ledger_ids) < 35:
        reporter.fail("ledger", f"too few ledger sources: {len(context.ledger_ids)}")
    else:
        reporter.passed("ledger", f"{len(context.ledger_ids)} unique source records")

    missing_sources = sorted(context.citations - context.ledger_ids)
    orphan_sources = sorted(context.ledger_ids - context.citations)
    if missing_sources:
        reporter.fail("citations", f"citations absent from ledger: {', '.join(missing_sources)}")
    else:
        reporter.passed("citations", f"all {len(context.citations)} cited source IDs resolve")
    if orphan_sources:
        reporter.warn("citations", f"ledger sources not cited by the manuscript: {', '.join(orphan_sources)}")
    if not context.citations:
        reporter.fail("citations", "manuscript contains no [Sxx] source citations")

    ledger_cutoff = extract_cutoff(ledger_text)
    if not ledger_cutoff:
        reporter.fail("ledger", "ledger verification date is absent")
    elif context.cutoff_date and ledger_cutoff != context.cutoff_date:
        reporter.fail("ledger", f"ledger date {ledger_cutoff} differs from book cutoff {context.cutoff_date}")
    else:
        reporter.passed("ledger", f"verification date agrees with cutoff: {ledger_cutoff}")
    return context


def required_output_strings(context: ReleaseContext) -> list[str]:
    values = [context.title, "Деньги", "Fraud journey", "Bot management", "Глоссарий"]
    if context.cutoff_date:
        values.append(context.cutoff_date)
    return [value for value in values if value]


def validate_pdf(path: Path, context: ReleaseContext, reporter: Reporter, min_pages: int, min_ratio: float) -> None:
    scope = "pdf"
    if not path.is_file():
        reporter.fail(scope, f"file not found: {path}")
        return
    try:
        import fitz
    except (ImportError, OSError) as exc:
        reporter.fail(scope, f"PyMuPDF is unavailable: {exc}")
        return

    try:
        document = fitz.open(path)
    except Exception as exc:
        reporter.fail(scope, f"cannot open PDF: {exc}")
        return
    try:
        if document.needs_pass:
            reporter.fail(scope, "PDF is encrypted or password protected")
            return
        page_count = len(document)
        if page_count < min_pages:
            reporter.fail(scope, f"too few pages: {page_count} < {min_pages}")
        else:
            reporter.passed(scope, f"page count: {page_count}")

        bad_sizes: list[str] = []
        page_texts: list[str] = []
        links: list[tuple[int, dict]] = []
        for number, page in enumerate(document, 1):
            dimensions = sorted((float(page.rect.width), float(page.rect.height)))
            target = sorted(A5_POINTS)
            if any(abs(actual - expected) > 3.0 for actual, expected in zip(dimensions, target)):
                bad_sizes.append(f"{number}:{page.rect.width:.1f}x{page.rect.height:.1f}pt")
            try:
                page_texts.append(page.get_text("text").strip())
                links.extend((number, link) for link in page.get_links())
            except Exception as exc:
                reporter.fail(scope, f"cannot inspect page {number}: {exc}")
                page_texts.append("")
        if bad_sizes:
            reporter.fail(scope, f"non-A5 pages: {', '.join(bad_sizes[:8])}")
        else:
            reporter.passed(scope, "all pages are A5 within 3 pt tolerance")

        blank_pages = [str(i) for i, text in enumerate(page_texts, 1) if len(re.sub(r"\s+", "", text)) < 12]
        if blank_pages:
            reporter.fail(scope, f"blank or nearly blank pages: {', '.join(blank_pages[:20])}")
        else:
            reporter.passed(scope, "no blank pages detected by text extraction")
        full_text = "\n".join(page_texts)
        ratio = len(re.sub(r"\s+", "", full_text)) / max(1, len(re.sub(r"\s+", "", context.plain_text)))
        if ratio < min_ratio:
            reporter.fail(scope, f"extracted text volume ratio is {ratio:.2f}, expected >= {min_ratio:.2f}")
        else:
            reporter.passed(scope, f"text volume ratio: {ratio:.2f}")
        missing = [value for value in required_output_strings(context) if value.casefold() not in full_text.casefold()]
        if missing:
            reporter.fail(scope, f"required text missing: {', '.join(missing)}")
        else:
            reporter.passed(scope, "required title/section/date strings are extractable")
        if "\ufffd" in full_text or "\u25a1" in full_text:
            reporter.fail(scope, "replacement or empty-square glyphs found in extracted text")

        metadata = document.metadata or {}
        title = (metadata.get("title") or "").strip()
        author = (metadata.get("author") or "").strip()
        if not title or normalize_visible(context.title) not in normalize_visible(title):
            reporter.fail(scope, f"PDF title metadata does not match manuscript: {title!r}")
        else:
            reporter.passed(scope, "title metadata matches")
        if context.author and context.author.casefold() not in author.casefold():
            reporter.fail(scope, f"PDF author metadata does not match manuscript: {author!r}")
        else:
            reporter.passed(scope, "author metadata matches")

        try:
            toc = document.get_toc(simple=True)
        except Exception as exc:
            reporter.fail(scope, f"cannot read PDF outline: {exc}")
            toc = []
        minimum_toc = min(10, max(1, len(context.h1_titles)))
        if len(toc) < minimum_toc:
            reporter.fail(scope, f"PDF outline has {len(toc)} entries; expected at least {minimum_toc}")
        elif not any("глоссарий" in str(item[1]).casefold() for item in toc if len(item) > 1):
            reporter.fail(scope, "PDF outline does not contain the glossary")
        else:
            reporter.passed(scope, f"PDF outline entries: {len(toc)}")

        unsafe_links: list[str] = []
        invalid_internal: list[str] = []
        for page_number, link in links:
            uri = str(link.get("uri") or "")
            if uri:
                scheme = urlsplit(uri).scheme.casefold()
                if scheme not in {"https", "http", "mailto"}:
                    unsafe_links.append(f"p{page_number}:{uri}")
                elif scheme == "http":
                    reporter.warn(scope, f"unencrypted external link on page {page_number}: {uri}")
            destination = link.get("page")
            if isinstance(destination, int) and destination >= page_count:
                invalid_internal.append(f"p{page_number}->page {destination + 1}")
        if unsafe_links:
            reporter.fail(scope, f"unsafe link schemes: {', '.join(unsafe_links[:8])}")
        if invalid_internal:
            reporter.fail(scope, f"invalid internal links: {', '.join(invalid_internal[:8])}")
        minimum_links = min(10, max(1, len(context.citations)))
        if len(links) < minimum_links:
            reporter.fail(scope, f"only {len(links)} PDF links; expected at least {minimum_links}")
        else:
            reporter.passed(scope, f"PDF links: {len(links)}")
    finally:
        document.close()


def safe_zip_names(names: Iterable[str]) -> list[str]:
    bad: list[str] = []
    for name in names:
        normalized = posixpath.normpath(name.replace("\\", "/"))
        if normalized.startswith("../") or normalized == ".." or normalized.startswith("/"):
            bad.append(name)
    return bad


def validate_docx(path: Path, context: ReleaseContext, reporter: Reporter, min_ratio: float) -> None:
    scope = "docx"
    if not path.is_file():
        reporter.fail(scope, f"file not found: {path}")
        return
    required_members = {"[Content_Types].xml", "_rels/.rels", "word/document.xml"}
    xml_roots: dict[str, ET.Element] = {}
    names: set[str] = set()
    try:
        with zipfile.ZipFile(path) as archive:
            names = set(archive.namelist())
            bad_name = archive.testzip()
            if bad_name:
                reporter.fail(scope, f"ZIP CRC failure: {bad_name}")
            unsafe = safe_zip_names(names)
            if unsafe:
                reporter.fail(scope, f"unsafe OOXML member paths: {', '.join(unsafe[:8])}")
            missing_members = sorted(required_members - names)
            if missing_members:
                reporter.fail(scope, f"missing OOXML members: {', '.join(missing_members)}")
                return
            for name in names:
                if name.endswith(".xml") or name.endswith(".rels"):
                    try:
                        xml_roots[name] = ET.fromstring(archive.read(name))
                    except (ET.ParseError, KeyError, OSError) as exc:
                        reporter.fail(scope, f"invalid XML in {name}: {exc}")
            document_xml = archive.read("word/document.xml").decode("utf-8", "replace")
    except (OSError, zipfile.BadZipFile) as exc:
        reporter.fail(scope, f"not a valid DOCX ZIP package: {exc}")
        return

    if "word/document.xml" not in xml_roots:
        reporter.fail(scope, "word/document.xml could not be parsed")
        return
    reporter.passed(scope, f"OOXML package and {len(xml_roots)} XML parts parse")

    try:
        from docx import Document
    except (ImportError, OSError) as exc:
        reporter.fail(scope, f"python-docx is unavailable: {exc}")
        return
    try:
        document = Document(path)
    except Exception as exc:
        reporter.fail(scope, f"python-docx cannot open the document: {exc}")
        return

    paragraphs = [paragraph.text for paragraph in document.paragraphs]
    table_text = [cell.text for table in document.tables for row in table.rows for cell in row.cells]
    full_text = "\n".join(paragraphs + table_text)
    compact_len = len(re.sub(r"\s+", "", full_text))
    expected_len = max(1, len(re.sub(r"\s+", "", context.plain_text)))
    ratio = compact_len / expected_len
    if ratio < min_ratio:
        reporter.fail(scope, f"text volume ratio is {ratio:.2f}, expected >= {min_ratio:.2f}")
    else:
        reporter.passed(scope, f"text volume ratio: {ratio:.2f}")
    missing = [value for value in required_output_strings(context) if value.casefold() not in full_text.casefold()]
    if missing:
        reporter.fail(scope, f"required text missing: {', '.join(missing)}")

    headings = [paragraph.text for paragraph in document.paragraphs if paragraph.style and paragraph.style.name.startswith("Heading ")]
    minimum_headings = max(12, len(context.h1_titles))
    if len(headings) < minimum_headings:
        reporter.fail(scope, f"only {len(headings)} Word heading paragraphs; expected at least {minimum_headings}")
    elif not any("глоссарий" in heading.casefold() for heading in headings):
        reporter.fail(scope, "Word headings do not include the glossary")
    else:
        reporter.passed(scope, f"Word headings: {len(headings)}")
    if len(document.tables) < context.expected_tables:
        reporter.fail(scope, f"only {len(document.tables)} tables; Markdown has {context.expected_tables}")
    else:
        reporter.passed(scope, f"Word tables: {len(document.tables)}")

    field_parts = [
        data
        for name, data in ((name, ET.tostring(root, encoding="unicode")) for name, root in xml_roots.items())
        if name.startswith("word/")
    ]
    field_xml = "\n".join(field_parts)
    if not re.search(r"\bTOC\b", field_xml, re.I):
        reporter.fail(scope, "TOC field is absent")
    else:
        reporter.passed(scope, "TOC field is present")
    if not re.search(r"\bPAGE\b", field_xml, re.I):
        reporter.fail(scope, "PAGE field is absent from document/header/footer XML")
    else:
        reporter.passed(scope, "PAGE field is present")

    root = xml_roots["word/document.xml"]
    ns = {"w": OOXML_NS, "r": OOXML_REL_NS}
    hyperlink_elements = root.findall(".//w:hyperlink", ns)
    bookmarks = {
        element.get(f"{{{OOXML_NS}}}name", "")
        for element in root.findall(".//w:bookmarkStart", ns)
        if element.get(f"{{{OOXML_NS}}}name", "") and not element.get(f"{{{OOXML_NS}}}name", "").startswith("_")
    }
    if not hyperlink_elements:
        reporter.fail(scope, "no Word hyperlinks found")
    else:
        reporter.passed(scope, f"Word hyperlinks: {len(hyperlink_elements)}")
    if not bookmarks:
        reporter.fail(scope, "no usable Word bookmarks found")
    else:
        reporter.passed(scope, f"Word bookmarks: {len(bookmarks)}")

    relationships: dict[str, tuple[str, str]] = {}
    rel_root = xml_roots.get("word/_rels/document.xml.rels")
    if rel_root is not None:
        for relation in rel_root.findall(f"{{{PKG_REL_NS}}}Relationship"):
            relationships[relation.get("Id", "")] = (
                relation.get("Target", ""),
                relation.get("TargetMode", ""),
            )
    for hyperlink in hyperlink_elements:
        relation_id = hyperlink.get(f"{{{OOXML_REL_NS}}}id")
        anchor = hyperlink.get(f"{{{OOXML_NS}}}anchor")
        if relation_id and relation_id not in relationships:
            reporter.fail(scope, f"hyperlink references missing relationship {relation_id}")
        if anchor and anchor not in bookmarks:
            reporter.fail(scope, f"hyperlink references missing bookmark {anchor}")

    if contains_raw_mermaid(full_text, document_xml):
        reporter.fail(scope, "raw Mermaid source is present instead of rendered diagrams")
    else:
        reporter.passed(scope, "no raw Mermaid source")
    if RAW_TEX_RE.search(full_text):
        reporter.fail(scope, "raw TeX commands are present")

    properties = document.core_properties
    metadata_ok = True
    if context.title.casefold() not in (properties.title or "").casefold():
        reporter.fail(scope, f"core title does not match: {properties.title!r}")
        metadata_ok = False
    if context.author and context.author.casefold() not in (properties.author or "").casefold():
        reporter.fail(scope, f"core author does not match: {properties.author!r}")
        metadata_ok = False
    if metadata_ok:
        reporter.passed(scope, "core title/author metadata match")

    bad_sections: list[str] = []
    for index, section in enumerate(document.sections, 1):
        dimensions = sorted((float(section.page_width.pt), float(section.page_height.pt)))
        if any(abs(actual - expected) > 3.0 for actual, expected in zip(dimensions, sorted(A5_POINTS))):
            bad_sections.append(str(index))
    if bad_sections:
        reporter.fail(scope, f"non-A5 Word sections: {', '.join(bad_sections)}")
    else:
        reporter.passed(scope, "all Word sections are A5")
    reporter.warn(scope, "structural checks do not replace a full-page LibreOffice/Word visual render review")


def resolve_epub_member(base: str, href: str) -> str:
    decoded = unquote(href.split("#", 1)[0])
    return posixpath.normpath(posixpath.join(base, decoded))


def validate_epub(
    path: Path,
    context: ReleaseContext,
    reporter: Reporter,
    min_chapters: int,
    min_ratio: float,
    epubcheck: str | None,
) -> None:
    scope = "epub"
    if not path.is_file():
        reporter.fail(scope, f"file not found: {path}")
        return
    try:
        archive = zipfile.ZipFile(path)
    except (OSError, zipfile.BadZipFile) as exc:
        reporter.fail(scope, f"not a valid EPUB ZIP package: {exc}")
        return

    with archive:
        infos = archive.infolist()
        names = set(archive.namelist())
        if archive.testzip():
            reporter.fail(scope, "EPUB ZIP CRC failure")
        unsafe = safe_zip_names(names)
        if unsafe:
            reporter.fail(scope, f"unsafe EPUB member paths: {', '.join(unsafe[:8])}")
        if not infos or infos[0].filename != "mimetype":
            reporter.fail(scope, "mimetype must be the first ZIP member")
        elif infos[0].compress_type != zipfile.ZIP_STORED:
            reporter.fail(scope, "mimetype must be stored without compression")
        elif archive.read("mimetype") != b"application/epub+zip":
            reporter.fail(scope, "invalid EPUB mimetype value")
        else:
            reporter.passed(scope, "mimetype is first, uncompressed and correct")

        if "META-INF/container.xml" not in names:
            reporter.fail(scope, "META-INF/container.xml is absent")
            return
        try:
            container = ET.fromstring(archive.read("META-INF/container.xml"))
        except ET.ParseError as exc:
            reporter.fail(scope, f"invalid container.xml: {exc}")
            return
        rootfiles = container.findall(f".//{{{CONTAINER_NS}}}rootfile")
        if not rootfiles:
            reporter.fail(scope, "container.xml has no rootfile")
            return
        opf_path = rootfiles[0].get("full-path", "")
        if not opf_path or opf_path not in names:
            reporter.fail(scope, f"OPF rootfile not found: {opf_path!r}")
            return
        try:
            opf = ET.fromstring(archive.read(opf_path))
        except ET.ParseError as exc:
            reporter.fail(scope, f"invalid OPF XML: {exc}")
            return
        opf_ns = opf.tag.split("}", 1)[0].lstrip("{") if "}" in opf.tag else ""
        ns = {"opf": opf_ns, "dc": DC_NS}
        metadata = opf.find("opf:metadata", ns)
        manifest_node = opf.find("opf:manifest", ns)
        spine_node = opf.find("opf:spine", ns)
        if metadata is None or manifest_node is None or spine_node is None:
            reporter.fail(scope, "OPF metadata, manifest or spine is missing")
            return

        title = element_text(metadata.find("dc:title", ns))
        creator = element_text(metadata.find("dc:creator", ns))
        language = element_text(metadata.find("dc:language", ns))
        identifier = element_text(metadata.find("dc:identifier", ns))
        dates = [element_text(item) for item in metadata.findall("dc:date", ns)]
        metadata_ok = True
        if context.title.casefold() not in title.casefold():
            reporter.fail(scope, f"EPUB title does not match: {title!r}")
            metadata_ok = False
        if context.author and context.author.casefold() not in creator.casefold():
            reporter.fail(scope, f"EPUB creator does not match: {creator!r}")
            metadata_ok = False
        if not language.casefold().startswith("ru"):
            reporter.fail(scope, f"EPUB language is not Russian: {language!r}")
            metadata_ok = False
        if not identifier:
            reporter.fail(scope, "EPUB unique identifier is empty")
            metadata_ok = False
        if context.cutoff_date and not any(date.startswith(context.cutoff_date) for date in dates):
            reporter.fail(scope, f"EPUB dc:date does not contain cutoff {context.cutoff_date}")
            metadata_ok = False
        if metadata_ok:
            reporter.passed(scope, "EPUB metadata has title/creator/language/identifier/date")

        opf_dir = posixpath.dirname(opf_path)
        manifest: dict[str, dict[str, str]] = {}
        for item in manifest_node.findall("opf:item", ns):
            item_id = item.get("id", "")
            href = item.get("href", "")
            resolved = resolve_epub_member(opf_dir, href)
            manifest[item_id] = {
                "href": href,
                "path": resolved,
                "media": item.get("media-type", ""),
                "properties": item.get("properties", ""),
            }
            if resolved not in names:
                reporter.fail(scope, f"manifest target is missing: {resolved}")
        spine_ids = [item.get("idref", "") for item in spine_node.findall("opf:itemref", ns)]
        missing_spine = [item_id for item_id in spine_ids if item_id not in manifest]
        if missing_spine:
            reporter.fail(scope, f"spine idrefs absent from manifest: {', '.join(missing_spine)}")
        nav_items = [item for item in manifest.values() if "nav" in item["properties"].split()]
        if len(nav_items) != 1:
            reporter.fail(scope, f"expected exactly one nav item, found {len(nav_items)}")
            return
        nav_path = nav_items[0]["path"]

        xml_cache: dict[str, ET.Element] = {}
        text_parts: list[str] = []
        raw_mermaid_files: list[str] = []
        raw_tex_files: list[str] = []
        table_count = 0
        content_links: list[tuple[str, str]] = []
        chapter_titles: list[str] = []
        for item_id in spine_ids:
            item = manifest.get(item_id)
            if not item or item["media"] not in {"application/xhtml+xml", "text/html"}:
                continue
            member = item["path"]
            try:
                root = ET.fromstring(archive.read(member))
            except (ET.ParseError, KeyError) as exc:
                reporter.fail(scope, f"invalid spine XHTML {member}: {exc}")
                continue
            xml_cache[member] = root
            text = element_text(root)
            text_parts.append(text)
            table_count += sum(1 for node in root.iter() if local_name(node.tag) == "table")
            first_h1 = next((node for node in root.iter() if local_name(node.tag) == "h1"), None)
            html_title = next((node for node in root.iter() if local_name(node.tag) == "title"), None)
            display_title = element_text(first_h1) or element_text(html_title)
            chapter_titles.append(display_title)
            if re.fullmatch(r"Раздел\s+\d+", element_text(html_title), re.I) or re.fullmatch(
                r"Раздел\s+\d+", display_title, re.I
            ):
                reporter.fail(scope, f"generic chapter title in {member}: {display_title!r}")
            serialized = ET.tostring(root, encoding="unicode")
            if contains_raw_mermaid(text, serialized):
                raw_mermaid_files.append(member)
            if RAW_TEX_RE.search(text):
                raw_tex_files.append(member)
            for node in root.iter():
                if local_name(node.tag) == "a" and node.get("href"):
                    content_links.append((member, node.get("href", "")))

        if len(chapter_titles) < min_chapters:
            reporter.fail(scope, f"only {len(chapter_titles)} readable spine chapters; expected at least {min_chapters}")
        elif not any("глоссарий" in title.casefold() for title in chapter_titles):
            reporter.fail(scope, "EPUB chapter titles do not include the glossary")
        else:
            reporter.passed(scope, f"readable spine chapters: {len(chapter_titles)}")
        if raw_mermaid_files:
            reporter.fail(scope, f"raw Mermaid source in: {', '.join(raw_mermaid_files[:8])}")
        else:
            reporter.passed(scope, "no raw Mermaid source")
        if raw_tex_files:
            reporter.fail(scope, f"raw TeX source in: {', '.join(raw_tex_files[:8])}")
        if table_count < context.expected_tables:
            reporter.fail(scope, f"only {table_count} EPUB tables; Markdown has {context.expected_tables}")
        else:
            reporter.passed(scope, f"EPUB tables: {table_count}")
        full_text = "\n".join(text_parts)
        ratio = len(re.sub(r"\s+", "", full_text)) / max(1, len(re.sub(r"\s+", "", context.plain_text)))
        if ratio < min_ratio:
            reporter.fail(scope, f"text volume ratio is {ratio:.2f}, expected >= {min_ratio:.2f}")
        else:
            reporter.passed(scope, f"text volume ratio: {ratio:.2f}")
        missing = [value for value in required_output_strings(context) if value.casefold() not in full_text.casefold()]
        if missing:
            reporter.fail(scope, f"required text missing: {', '.join(missing)}")

        try:
            nav_root = ET.fromstring(archive.read(nav_path))
        except (ET.ParseError, KeyError) as exc:
            reporter.fail(scope, f"invalid nav XHTML: {exc}")
            nav_root = None
        nav_links: list[tuple[str, str]] = []
        if nav_root is not None:
            nav_texts = []
            for node in nav_root.iter():
                if local_name(node.tag) == "a" and node.get("href"):
                    nav_links.append((nav_path, node.get("href", "")))
                    nav_texts.append(element_text(node))
            generic = [text for text in nav_texts if re.fullmatch(r"Раздел\s+\d+", text, re.I)]
            if generic:
                reporter.fail(scope, f"generic nav labels: {', '.join(generic[:8])}")
            if len(nav_links) < min_chapters:
                reporter.fail(scope, f"nav has only {len(nav_links)} links")
            else:
                reporter.passed(scope, f"nav links: {len(nav_links)}")

        all_links = content_links + nav_links
        external_links = 0
        bad_links: list[str] = []
        id_cache: dict[str, set[str]] = {}
        for source_member, href in all_links:
            parsed = urlsplit(href)
            if parsed.scheme:
                if parsed.scheme.casefold() not in {"https", "http", "mailto"}:
                    bad_links.append(f"{source_member}:{href}")
                elif parsed.scheme.casefold() == "http":
                    reporter.warn(scope, f"unencrypted external link: {href}")
                else:
                    external_links += 1
                continue
            target_member = resolve_epub_member(posixpath.dirname(source_member), href)
            if not href.split("#", 1)[0]:
                target_member = source_member
            if target_member not in names:
                bad_links.append(f"{source_member}:{href}")
                continue
            fragment = parsed.fragment
            if fragment:
                if target_member not in id_cache:
                    try:
                        target_root = xml_cache.get(target_member)
                        if target_root is None:
                            target_root = ET.fromstring(archive.read(target_member))
                        id_cache[target_member] = {
                            node.get("id", "") for node in target_root.iter() if node.get("id")
                        }
                    except (ET.ParseError, KeyError):
                        id_cache[target_member] = set()
                if unquote(fragment) not in id_cache[target_member]:
                    bad_links.append(f"{source_member}:{href}")
        if bad_links:
            reporter.fail(scope, f"broken or unsafe EPUB links: {', '.join(bad_links[:12])}")
        non_nav_links = [item for item in content_links if item[0] != nav_path]
        if not non_nav_links or external_links == 0:
            reporter.fail(scope, "EPUB body has no usable external source links")
        else:
            reporter.passed(scope, f"EPUB content links: {len(non_nav_links)}, external: {external_links}")

    command = epubcheck or shutil.which("epubcheck")
    if command:
        try:
            result = subprocess.run(
                [command, str(path)],
                capture_output=True,
                text=True,
                timeout=180,
                check=False,
            )
        except (OSError, subprocess.SubprocessError) as exc:
            reporter.fail(scope, f"epubcheck could not run: {exc}")
        else:
            if result.returncode:
                detail = (result.stdout + "\n" + result.stderr).strip().replace("\n", " | ")
                reporter.fail(scope, f"epubcheck failed ({result.returncode}): {detail[:800]}")
            else:
                reporter.passed(scope, "epubcheck passed")
    else:
        reporter.warn(scope, "epubcheck not found; structural EPUB checks ran, official validation did not")


def path_within(path: Path, parent: Path) -> bool:
    try:
        path.resolve().relative_to(parent.resolve())
        return True
    except (OSError, ValueError):
        return False


def manifest_input_records(data: dict) -> list[dict]:
    """Accept list records and both role-keyed and path-keyed input mappings."""
    records: list[dict] = []
    inputs = data.get("inputs")
    if isinstance(inputs, list):
        records.extend(dict(item) for item in inputs if isinstance(item, dict))
    elif isinstance(inputs, dict):
        for key, value in inputs.items():
            if isinstance(value, str):
                if key in REQUIRED_MANIFEST_INPUT_ROLES:
                    records.append({"role": key, "path": value})
                else:
                    records.append({"path": key, "sha256": value})
            elif isinstance(value, dict):
                record = dict(value)
                if key in REQUIRED_MANIFEST_INPUT_ROLES:
                    record.setdefault("role", key)
                elif not record.get("path") and not record.get("file"):
                    record["path"] = key
                records.append(record)
    if records:
        return records

    legacy_fields = (
        ("source", "source", "source_sha256"),
        ("source_ledger", "source_ledger", "source_ledger_sha256"),
        ("ledger", "source_ledger", "ledger_sha256"),
        ("claim_map", "claim_map", "claim_map_sha256"),
        ("builder", "builder", "builder_sha256"),
        ("requirements", "requirements", "requirements_sha256"),
        ("package_json", "package_json", "package_json_sha256"),
        ("package_lock", "package_lock", "package_lock_sha256"),
        ("puppeteer_config", "puppeteer_config", "puppeteer_config_sha256"),
    )
    for field_name, role, hash_field in legacy_fields:
        value = data.get(field_name)
        if isinstance(value, dict):
            record = dict(value)
            record.setdefault("role", role)
            records.append(record)
        elif isinstance(value, str):
            records.append(
                {
                    "role": role,
                    "path": value,
                    "sha256": data.get(hash_field),
                }
            )
    return records


def validate_manifest(
    path: Path,
    root: Path,
    context: ReleaseContext,
    artifacts: dict[str, Path],
    claim_map: Path,
    reporter: Reporter,
) -> None:
    scope = "manifest"
    text = read_utf8(path, reporter, scope)
    if text is None:
        return
    try:
        data = json.loads(text)
    except json.JSONDecodeError as exc:
        reporter.fail(scope, f"invalid JSON: {exc}")
        return
    if not isinstance(data, dict):
        reporter.fail(scope, "manifest root must be an object")
        return

    input_records = manifest_input_records(data)
    raw_artifacts = data.get("artifacts")
    artifact_records: list[dict] = []
    if isinstance(raw_artifacts, list):
        artifact_records = [dict(item) for item in raw_artifacts if isinstance(item, dict)]
    elif isinstance(raw_artifacts, dict):
        for key, value in raw_artifacts.items():
            if isinstance(value, str):
                artifact_records.append({"path": key, "sha256": value})
            elif isinstance(value, dict):
                record = dict(value)
                if not record.get("path") and not record.get("file"):
                    record["path"] = key
                artifact_records.append(record)
    seen_paths: set[Path] = set()
    verified_paths: set[Path] = set()
    verified_roles: dict[str, Path] = {}

    def verify_record(
        record: dict,
        base: Path,
        record_kind: str,
        *,
        allow_external_absolute: bool,
    ) -> tuple[Path | None, bool]:
        raw_path = record.get("path") or record.get("file")
        expected_hash = record.get("sha256")
        if not isinstance(raw_path, str) or not raw_path:
            reporter.fail(scope, f"{record_kind} record has no file/path")
            return None, False
        alternate_path = record.get("file") if record.get("path") is not None else None
        if isinstance(alternate_path, str) and alternate_path != raw_path:
            reporter.fail(scope, f"{record_kind} record has conflicting path and file values")
            return None, False
        candidate = Path(raw_path)
        resolved = candidate.resolve() if candidate.is_absolute() else (base / candidate).resolve()
        if (not candidate.is_absolute() or not allow_external_absolute) and not path_within(
            resolved, root
        ):
            reporter.fail(scope, f"{record_kind} path escapes release root: {raw_path}")
            return None, False
        if resolved in seen_paths:
            reporter.fail(scope, f"duplicate manifest path: {raw_path}")
        seen_paths.add(resolved)
        if not resolved.is_file():
            reporter.fail(scope, f"manifest file is missing: {raw_path}")
            return resolved, False
        if not isinstance(expected_hash, str) or not re.fullmatch(r"[0-9a-fA-F]{64}", expected_hash):
            reporter.fail(scope, f"{record_kind} has no valid SHA-256: {raw_path}")
            return resolved, False
        actual_hash = sha256(resolved)
        hash_ok = actual_hash.casefold() == expected_hash.casefold()
        if not hash_ok:
            reporter.fail(scope, f"SHA-256 mismatch: {raw_path}")
        else:
            verified_paths.add(resolved)
        expected_bytes = record.get("bytes")
        if expected_bytes is not None and expected_bytes != resolved.stat().st_size:
            reporter.fail(scope, f"byte-size mismatch: {raw_path}")
            hash_ok = False
        return resolved, hash_ok

    expected_role_paths: dict[str, set[Path]] = {
        "source": {context.book_path.resolve()},
        "source_ledger": {context.ledger_path.resolve()},
        "claim_map": {claim_map.resolve()},
        "builder": {
            (root / "build_release.py").resolve(),
            (root / "build_release_local.py").resolve(),
        },
        "requirements": {(root / "requirements-build.txt").resolve()},
        "package_json": {(root / "package.json").resolve()},
        "package_lock": {(root / "package-lock.json").resolve()},
    }

    def infer_input_role(resolved: Path | None) -> str:
        if resolved is None:
            return ""
        for role, expected in expected_role_paths.items():
            if resolved in expected:
                return role
        if "puppeteer" in resolved.name.casefold() and resolved.suffix.casefold() == ".json":
            return "puppeteer_config"
        return ""

    for record in input_records:
        resolved, hash_ok = verify_record(
            record,
            root,
            "input",
            allow_external_absolute=True,
        )
        declared_role = record.get("role")
        role = declared_role.strip() if isinstance(declared_role, str) else ""
        role = role or infer_input_role(resolved)
        if not role:
            reporter.fail(scope, f"input record has no recognizable role: {record!r}")
            continue
        if role in verified_roles:
            reporter.fail(scope, f"duplicate manifest input role: {role}")
            continue
        if hash_ok and resolved is not None:
            verified_roles[role] = resolved
    for record in artifact_records:
        verify_record(
            record,
            path.parent,
            "artifact",
            allow_external_absolute=False,
        )

    missing_roles = [role for role in REQUIRED_MANIFEST_INPUT_ROLES if role not in verified_roles]
    if missing_roles:
        reporter.fail(scope, f"required input hashes absent or invalid: {', '.join(missing_roles)}")
    binding_failures = 0
    for role, expected_paths in expected_role_paths.items():
        resolved = verified_roles.get(role)
        if resolved is not None and resolved not in expected_paths:
            binding_failures += 1
            reporter.fail(scope, f"input role {role} points to the wrong file: {resolved}")
    puppeteer_path = verified_roles.get("puppeteer_config")
    if puppeteer_path is not None and (
        "puppeteer" not in puppeteer_path.name.casefold()
        or puppeteer_path.suffix.casefold() != ".json"
    ):
        binding_failures += 1
        reporter.fail(scope, f"input role puppeteer_config is not a Puppeteer JSON file: {puppeteer_path}")
    if not missing_roles and not binding_failures:
        reporter.passed(scope, "all required build inputs are role-bound and hash-verified")

    missing_artifacts = [name for name, artifact in artifacts.items() if artifact.resolve() not in verified_paths]
    if missing_artifacts:
        reporter.fail(scope, f"artifact hashes absent or invalid: {', '.join(missing_artifacts)}")
    else:
        reporter.passed(scope, "PDF, DOCX and EPUB hashes are verified")

    manifest_cutoff = data.get("cutoff_date") or data.get("source_date")
    if not isinstance(manifest_cutoff, str) or not re.fullmatch(r"20\d{2}-\d{2}-\d{2}", manifest_cutoff):
        reporter.fail(scope, "manifest has no valid explicit cutoff_date/source_date")
    elif context.cutoff_date and manifest_cutoff != context.cutoff_date:
        reporter.fail(scope, f"manifest cutoff {manifest_cutoff!r} differs from {context.cutoff_date}")
    else:
        reporter.passed(scope, "manifest cutoff agrees with manuscript")

    schema_version = data.get("schema_version")
    if not isinstance(schema_version, (str, int, float)) or not str(schema_version).strip():
        reporter.fail(scope, "manifest has no schema_version")
    else:
        reporter.passed(scope, f"manifest schema_version is declared: {schema_version}")

    tools = data.get("tools")
    tools_ok = isinstance(tools, dict)
    if not tools_ok:
        reporter.fail(scope, "manifest has no tools version map")
        tools = {}
    missing_tools = [name for name in REQUIRED_MANIFEST_TOOLS if name not in tools]
    if missing_tools:
        tools_ok = False
        reporter.fail(scope, f"manifest tool versions are missing: {', '.join(missing_tools)}")

    python_value = tools.get("python")
    python_version = (
        python_value.get("version") if isinstance(python_value, dict) else python_value
    )
    if not isinstance(python_version, str) or not python_version.strip():
        tools_ok = False
        reporter.fail(scope, "manifest Python version is empty")

    package_versions = tools.get("python_packages")
    if not isinstance(package_versions, dict) or not package_versions or any(
        not isinstance(version, str) or not version.strip()
        for version in package_versions.values()
    ):
        tools_ok = False
        reporter.fail(scope, "manifest python_packages must be a non-empty version map")

    for tool_name in ("weasyprint", "mmdc"):
        tool_record = tools.get(tool_name)
        if not isinstance(tool_record, dict):
            tools_ok = False
            reporter.fail(scope, f"manifest {tool_name} version/hash record is missing")
            continue
        version = tool_record.get("version")
        raw_tool_path = tool_record.get("path") or tool_record.get("file")
        expected_hash = tool_record.get("sha256")
        if not isinstance(version, str) or not version.strip():
            tools_ok = False
            reporter.fail(scope, f"manifest {tool_name} version is empty")
        if not isinstance(raw_tool_path, str) or not raw_tool_path:
            tools_ok = False
            reporter.fail(scope, f"manifest {tool_name} path is empty")
            continue
        tool_candidate = Path(raw_tool_path)
        tool_path = (
            tool_candidate.resolve()
            if tool_candidate.is_absolute()
            else (root / tool_candidate).resolve()
        )
        if not tool_path.is_file():
            tools_ok = False
            reporter.fail(scope, f"manifest {tool_name} executable is missing: {raw_tool_path}")
            continue
        if not isinstance(expected_hash, str) or not re.fullmatch(
            r"[0-9a-fA-F]{64}", expected_hash
        ):
            tools_ok = False
            reporter.fail(scope, f"manifest {tool_name} has no valid SHA-256")
        elif sha256(tool_path).casefold() != expected_hash.casefold():
            tools_ok = False
            reporter.fail(scope, f"manifest {tool_name} SHA-256 mismatch")
    if tools_ok:
        reporter.passed(scope, "required tool versions and executable hashes are verified")


def resolve_argument(root: Path, value: str | None, default: Path) -> Path:
    if value is None:
        return default
    candidate = Path(value)
    return candidate.resolve() if candidate.is_absolute() else (root / candidate).resolve()


def run_self_test() -> int:
    try:
        import fitz
    except (ImportError, OSError) as exc:
        print(f"SELF-TEST FAIL: PyMuPDF is required to create the blank-PDF fixture: {exc}")
        return 1

    results: list[tuple[str, bool, str]] = []
    short_words = count_unicode_words("слово " * 5_000)
    release_words = count_unicode_words("слово " * 16_000)
    volume_gate_ok = short_words < 16_000 <= release_words
    results.append(
        (
            "minimum volume counter",
            volume_gate_ok,
            f"M1={short_words}, M2={release_words}, minimum=16000",
        )
    )
    heading_contract_ok = (
        normalize_heading("28. Глоссарий") == normalize_heading("Глоссарий")
        and normalize_heading("1. Деньги, риск") == "деньги, риск"
    )
    results.append(
        ("semantic heading normalization", heading_contract_ok, "leading numbers ignored")
    )
    mermaid_contract_ok = (
        not contains_raw_mermaid(
            "detect --> triage --> contain\n---\nGraph — типизированная сеть."
        )
        and contains_raw_mermaid("flowchart LR\nA --> B")
        and contains_raw_mermaid("rendered", '<code class="language-mermaid">')
    )
    results.append(
        ("Mermaid declaration detection", mermaid_contract_ok, "ordinary arrows are ignored")
    )
    title_contract_ok = normalize_visible("АНТИФРОД\u00a0 2026") == normalize_visible(
        "АНТИФРОД 2026"
    )
    results.append(("visible title normalization", title_contract_ok, "Unicode spaces collapse"))
    parsed_role_records = manifest_input_records(
        {
            "inputs": {
                "source": {"path": "output/book.md", "sha256": "0" * 64},
            }
        }
    )
    role_path_contract_ok = (
        len(parsed_role_records) == 1
        and parsed_role_records[0].get("role") == "source"
        and parsed_role_records[0].get("path") == "output/book.md"
    )
    results.append(
        ("manifest role/path schema", role_path_contract_ok, "role-keyed mapping accepted")
    )
    with tempfile.TemporaryDirectory(prefix="antifraud_release_gate_") as temp_name:
        root = Path(temp_name)
        (root / "output").mkdir()
        (root / "research").mkdir()
        (root / "dist").mkdir()
        book = root / "output" / f"{BASE_NAME}.md"
        ledger = root / "research" / "SOURCE_LEDGER.md"
        claim_map = root / "research" / "CLAIM_MAP.md"
        book.write_text(
            "---\ntitle: Test\nauthor: Test\ndate: 2026-07-21\n---\n\n# Test\nMissing source [S99].\n",
            encoding="utf-8",
        )
        ledger.write_text(
            "Дата проверки: 2026-07-21\n\n| ID | Org | Doc | URL |\n|---|---|---|---|\n"
            "| S01 | Test | Test | https://example.invalid/source |\n",
            encoding="utf-8",
        )
        claim_map.write_text("test\n", encoding="utf-8")

        markdown_report = Reporter(quiet=True)
        context = validate_markdown_and_ledger(book, ledger, markdown_report)
        results.append(
            (
                "missing citation/structure",
                markdown_report.failures > 0,
                f"{markdown_report.failures} expected failures",
            )
        )

        pdf_path = root / "dist" / f"{BASE_NAME}.pdf"
        blank_pdf = fitz.open()
        for _ in range(20):
            blank_pdf.new_page(width=A5_POINTS[0], height=A5_POINTS[1])
        blank_pdf.save(pdf_path)
        blank_pdf.close()
        pdf_report = Reporter(quiet=True)
        validate_pdf(pdf_path, context, pdf_report, min_pages=20, min_ratio=0.50)
        results.append(("blank PDF", pdf_report.failures > 0, f"{pdf_report.failures} expected failures"))

        docx_path = root / "dist" / f"{BASE_NAME}.docx"
        epub_path = root / "dist" / f"{BASE_NAME}.epub"
        for path in (docx_path, epub_path):
            with zipfile.ZipFile(path, "w") as archive:
                archive.writestr("junk.txt", "not a document")
        docx_report = Reporter(quiet=True)
        validate_docx(docx_path, context, docx_report, min_ratio=0.50)
        results.append(
            ("junk DOCX ZIP", docx_report.failures > 0, f"{docx_report.failures} expected failures")
        )
        epub_report = Reporter(quiet=True)
        validate_epub(epub_path, context, epub_report, 15, 0.50, epubcheck=None)
        results.append(
            ("junk EPUB ZIP", epub_report.failures > 0, f"{epub_report.failures} expected failures")
        )

        builder = root / "build_release_local.py"
        requirements = root / "requirements-build.txt"
        package_json = root / "package.json"
        package_lock = root / "package-lock.json"
        puppeteer_config = root / "mermaid-puppeteer.json"
        tools_dir = root / "tools"
        tools_dir.mkdir()
        weasyprint_tool = tools_dir / "weasyprint.exe"
        mmdc_tool = tools_dir / "mmdc.cmd"
        for fixture_path, fixture_text in (
            (builder, "# builder\n"),
            (requirements, "fixture==1.0\n"),
            (package_json, "{}\n"),
            (package_lock, "{}\n"),
            (puppeteer_config, "{}\n"),
            (weasyprint_tool, "fixture\n"),
            (mmdc_tool, "fixture\n"),
        ):
            fixture_path.write_text(fixture_text, encoding="utf-8")

        manifest_path = root / "dist" / "manifest.json"
        role_paths = {
            "source": book,
            "source_ledger": ledger,
            "claim_map": claim_map,
            "builder": builder,
            "requirements": requirements,
            "package_json": package_json,
            "package_lock": package_lock,
            "puppeteer_config": puppeteer_config,
        }
        good_manifest = {
            "schema_version": "1.0",
            "cutoff_date": "2026-07-21",
            "inputs": {
                role: {
                    "path": fixture_path.relative_to(root).as_posix(),
                    "bytes": fixture_path.stat().st_size,
                    "sha256": sha256(fixture_path),
                }
                for role, fixture_path in role_paths.items()
            },
            "tools": {
                "python": "3.test",
                "python_packages": {"fixture": "1.0"},
                "weasyprint": {
                    "path": weasyprint_tool.relative_to(root).as_posix(),
                    "version": "test",
                    "sha256": sha256(weasyprint_tool),
                },
                "mmdc": {
                    "path": mmdc_tool.relative_to(root).as_posix(),
                    "version": "test",
                    "sha256": sha256(mmdc_tool),
                },
            },
            "artifacts": [
                {
                    "path": artifact.name,
                    "bytes": artifact.stat().st_size,
                    "sha256": sha256(artifact),
                }
                for artifact in (pdf_path, docx_path, epub_path)
            ],
        }
        manifest_path.write_text(json.dumps(good_manifest), encoding="utf-8")
        good_manifest_report = Reporter(quiet=True)
        validate_manifest(
            manifest_path,
            root,
            context,
            {"pdf": pdf_path, "docx": docx_path, "epub": epub_path},
            claim_map,
            good_manifest_report,
        )
        results.append(
            (
                "complete manifest contract",
                good_manifest_report.failures == 0,
                f"{good_manifest_report.failures} unexpected failures",
            )
        )

        bad_manifest = json.loads(json.dumps(good_manifest))
        bad_manifest["inputs"].pop("package_lock")
        bad_manifest["inputs"]["source"]["sha256"] = "0" * 64
        manifest_path.write_text(json.dumps(bad_manifest), encoding="utf-8")
        manifest_report = Reporter(quiet=True)
        validate_manifest(
            manifest_path,
            root,
            context,
            {"pdf": pdf_path, "docx": docx_path, "epub": epub_path},
            claim_map,
            manifest_report,
        )
        results.append(
            (
                "bad manifest hashes",
                manifest_report.failures > 0,
                f"{manifest_report.failures} expected failures",
            )
        )

    failed = False
    for name, check_passed, detail in results:
        if check_passed:
            print(f"SELF-TEST PASS [{name}] ({detail})")
        else:
            failed = True
            print(f"SELF-TEST FAIL [{name}] ({detail})")
    print(
        "SELF-TEST FAIL"
        if failed
        else "SELF-TEST PASS: all contracts held and all negative fixtures were rejected"
    )
    return 1 if failed else 0


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(
        description="Read-only local release gate for Markdown, sources, PDF, DOCX, EPUB and manifest.",
        formatter_class=argparse.ArgumentDefaultsHelpFormatter,
    )
    parser.add_argument(
        "root",
        nargs="?",
        default=str(Path(__file__).resolve().parents[1]),
        help="release root; relative artifact path options are resolved against it",
    )
    parser.add_argument("--book", help="canonical Markdown path")
    parser.add_argument("--ledger", help="source ledger path")
    parser.add_argument("--claim-map", help="claim map path")
    parser.add_argument("--pdf", help="PDF artifact path")
    parser.add_argument("--docx", help="DOCX artifact path")
    parser.add_argument("--epub", help="EPUB artifact path")
    parser.add_argument("--manifest", help="manifest JSON path")
    parser.add_argument("--epubcheck", help="explicit epubcheck executable; PATH is auto-detected otherwise")
    parser.add_argument("--min-pdf-pages", type=int, default=20)
    parser.add_argument("--min-epub-chapters", type=int, default=15)
    parser.add_argument("--min-text-ratio", type=float, default=0.50)
    parser.add_argument(
        "--min-book-words",
        type=int,
        default=16_000,
        help="minimum Unicode word count for the canonical manuscript",
    )
    parser.add_argument("--strict-warnings", action="store_true", help="treat WARN as a nonzero result")
    parser.add_argument("--self-test", action="store_true", help="run disposable negative-fixture tests only")
    return parser


def main(argv: list[str] | None = None) -> int:
    parser = build_parser()
    args = parser.parse_args(argv)
    if args.self_test:
        return run_self_test()
    if args.min_pdf_pages < 1 or args.min_epub_chapters < 1 or args.min_book_words < 1:
        parser.error("minimum page/chapter/word counts must be positive")
    if not 0.05 <= args.min_text_ratio <= 1.25:
        parser.error("--min-text-ratio must be between 0.05 and 1.25")

    root = Path(args.root).resolve()
    reporter = Reporter()
    if not root.is_dir():
        reporter.fail("root", f"release root not found: {root}")
        return reporter.summary(args.strict_warnings)
    reporter.passed("root", f"read-only validation root: {root}")

    book = resolve_argument(root, args.book, root / "output" / f"{BASE_NAME}.md")
    ledger = resolve_argument(root, args.ledger, root / "research" / "SOURCE_LEDGER.md")
    claim_map = resolve_argument(root, args.claim_map, root / "research" / "CLAIM_MAP.md")
    pdf = resolve_argument(root, args.pdf, root / "dist" / f"{BASE_NAME}.pdf")
    docx = resolve_argument(root, args.docx, root / "dist" / f"{BASE_NAME}.docx")
    epub = resolve_argument(root, args.epub, root / "dist" / f"{BASE_NAME}.epub")
    manifest = resolve_argument(root, args.manifest, root / "dist" / "manifest.json")

    context = validate_markdown_and_ledger(book, ledger, reporter, args.min_book_words)
    validate_pdf(pdf, context, reporter, args.min_pdf_pages, args.min_text_ratio)
    validate_docx(docx, context, reporter, args.min_text_ratio)
    validate_epub(
        epub,
        context,
        reporter,
        args.min_epub_chapters,
        args.min_text_ratio,
        args.epubcheck,
    )
    validate_manifest(
        manifest,
        root,
        context,
        {"pdf": pdf, "docx": docx, "epub": epub},
        claim_map,
        reporter,
    )
    return reporter.summary(args.strict_warnings)


if __name__ == "__main__":
    sys.exit(main())
