#!/usr/bin/env python3
"""Repeatably build A5 PDF, DOCX and EPUB from the canonical Markdown.

Binary output is deliberately written to ignored dist/. Run from any directory.
"""
from __future__ import annotations
import argparse, hashlib, html, json, re, shutil, subprocess, sys, zipfile
from pathlib import Path

ROOT = Path(__file__).resolve().parent
SOURCE = ROOT / "output" / "ANTIFRAUD_ULTIMATE_GUIDE_2026_RU.md"
DIST = ROOT / "dist"
BASENAME = "ANTIFRAUD_ULTIMATE_GUIDE_2026_RU"
ASSETS = DIST / "assets"

CSS = r"""
@page { size: A5; margin: 15mm 13mm 17mm; @bottom-center { content: counter(page); font: 8pt sans-serif; color: #667085 } }
@page:first { @bottom-center { content: none } }
html { font-family: "DejaVu Sans", sans-serif; font-size: 9.3pt; line-height: 1.36; color: #17212b }
body { hyphens: auto }
h1 { font-size: 19pt; color: #0b486b; break-before: page; margin: 0 0 8mm }
h1:first-of-type { break-before: avoid }
h2 { font-size: 13.5pt; color: #11698e; margin: 6mm 0 2mm; break-after: avoid }
h3 { font-size: 11pt; color: #174a65; break-after: avoid }
p { margin: 0 0 2.3mm; orphans: 3; widows: 3 }
a { color: #11698e; text-decoration: none }
table { width: 100%; border-collapse: collapse; font-size: 7.4pt; margin: 3mm 0; break-inside: avoid }
th { background: #dceef5; color: #173b4f }
th, td { border: .25mm solid #a9bec8; padding: 1.2mm; vertical-align: top }
tr { break-inside: avoid }
code { font-family: "DejaVu Sans Mono", monospace; font-size: .9em; background: #eef3f5 }
pre { white-space: pre-wrap; background: #eef3f5; border-left: 1mm solid #3d91b8; padding: 2.5mm; font-size: 7.2pt; break-inside: avoid }
blockquote { border-left: 1mm solid #d08b28; margin-left: 0; padding-left: 3mm }
ul, ol { padding-left: 5mm; margin-top: 1mm }
.title-page { height: 160mm; display: flex; flex-direction: column; justify-content: center; text-align: center }
.meta { color: #667085; font-size: 9pt }
"""

def require_modules():
    try:
        import markdown  # noqa
        import weasyprint  # noqa
        import docx  # noqa
        import ebooklib  # noqa
    except ImportError as exc:
        raise SystemExit(f"Missing build dependency: {exc}. Run: python -m pip install -r {ROOT/'requirements-build.txt'}")

def strip_frontmatter(text: str) -> tuple[dict[str, str], str]:
    meta: dict[str, str] = {}
    if text.startswith("---\n"):
        end = text.find("\n---\n", 4)
        if end >= 0:
            for line in text[4:end].splitlines():
                if ":" in line:
                    k, v = line.split(":", 1); meta[k.strip()] = v.strip().strip('"')
            return meta, text[end + 5:]
    return meta, text

def markdown_html(body: str) -> str:
    import markdown
    return markdown.markdown(body, extensions=["tables", "fenced_code", "sane_lists", "toc"], output_format="html5")

def render_mermaid(body: str) -> tuple[str, list[dict[str, Path | str]]]:
    """Render every Mermaid fence to SVG and PNG, retaining an accessible caption."""
    ASSETS.mkdir(parents=True, exist_ok=True)
    diagrams: list[dict[str, Path | str]] = []
    pattern = re.compile(r"```mermaid\n(.*?)\n```", re.S)
    def repl(match):
        idx = len(diagrams) + 1
        source = ASSETS / f"diagram_{idx:02}.mmd"
        svg = ASSETS / f"diagram_{idx:02}.svg"
        png = ASSETS / f"diagram_{idx:02}.png"
        caption = f"Схема {idx}: защитный поток и состояния"
        source.write_text(match.group(1) + "\n", encoding="utf-8")
        command=[str(ROOT/"node_modules/.bin/mmdc"), "-p", str(ROOT/"mermaid-puppeteer.json"), "-i", str(source)]
        subprocess.run(command+["-o", str(svg), "-b", "white"], check=True)
        subprocess.run(command+["-o", str(png), "-b", "white", "-s", "2"], check=True)
        diagrams.append({"source": source, "svg": svg, "png": png, "caption": caption})
        return f"![{caption}]({svg.as_posix()})\n\n*{caption}.*"
    return pattern.sub(repl, body), diagrams

def full_html(meta: dict[str, str], body_html: str) -> str:
    title = html.escape(meta.get("title", BASENAME))
    subtitle = html.escape(meta.get("subtitle", ""))
    date = html.escape(meta.get("date", ""))
    return f"""<!doctype html><html lang='ru'><head><meta charset='utf-8'><title>{title}</title><style>{CSS}</style></head>
<body><section class='title-page'><h1>{title}</h1><h2>{subtitle}</h2><p class='meta'>{date}</p></section>{body_html}</body></html>"""

def build_pdf(doc_html: str, target: Path):
    from weasyprint import HTML
    HTML(string=doc_html, base_url=str(ROOT)).write_pdf(target)

def add_inline(paragraph, text: str):
    # Preserve readable text while styling simple Markdown emphasis/code.
    parts = re.split(r"(`[^`]+`|\*\*[^*]+\*\*)", text)
    for part in parts:
        if part.startswith("`") and part.endswith("`"):
            run = paragraph.add_run(part[1:-1]); run.font.name = "DejaVu Sans Mono"
        elif part.startswith("**") and part.endswith("**"):
            paragraph.add_run(part[2:-2]).bold = True
        else:
            paragraph.add_run(part)

def build_docx(meta: dict[str, str], body: str, target: Path, diagrams):
    from docx import Document
    from docx.enum.section import WD_SECTION
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Mm, Pt
    doc = Document(); sec = doc.sections[0]
    sec.page_height, sec.page_width = Mm(210), Mm(148)
    sec.top_margin = sec.bottom_margin = Mm(15); sec.left_margin = sec.right_margin = Mm(14)
    normal = doc.styles["Normal"]; normal.font.name = "DejaVu Sans"; normal.font.size = Pt(9)
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(meta.get("title", BASENAME)); r.bold = True; r.font.size = Pt(24)
    p = doc.add_paragraph(meta.get("subtitle", "")); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p = doc.add_paragraph(meta.get("date", "")); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_page_break()
    doc.add_heading("Содержание", 1)
    # A real Word TOC field: Word/LibreOffice can update page numbers and links
    # after pagination instead of shipping a manually copied heading list.
    toc_paragraph = doc.add_paragraph()
    run = toc_paragraph.add_run()
    begin = OxmlElement("w:fldChar"); begin.set(qn("w:fldCharType"), "begin"); begin.set(qn("w:dirty"), "true")
    instruction = OxmlElement("w:instrText"); instruction.set(qn("xml:space"), "preserve"); instruction.text = ' TOC \\o "1-3" \\h \\z \\u '
    separate = OxmlElement("w:fldChar"); separate.set(qn("w:fldCharType"), "separate")
    placeholder = OxmlElement("w:t"); placeholder.text = "Обновите поле содержания при открытии документа."
    end = OxmlElement("w:fldChar"); end.set(qn("w:fldCharType"), "end")
    for node in (begin, instruction, separate, placeholder, end): run._r.append(node)
    doc.add_paragraph("Резервное содержание до обновления поля:")
    for title in re.findall(r"^# (.+)$", body, re.M):
        doc.add_paragraph(title, style="List Number")
    doc.add_page_break()
    lines = body.splitlines(); i = 0; in_code = False; code=[]
    while i < len(lines):
        line = lines[i]
        if line.startswith("```mermaid"):
            i += 1; mer=[]
            while i < len(lines) and not lines[i].startswith("```"): mer.append(lines[i]); i += 1
            d=diagrams.pop(0); p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
            p.add_run().add_picture(str(d["png"]), width=Mm(116)); doc.add_paragraph(str(d["caption"])).alignment=WD_ALIGN_PARAGRAPH.CENTER
            i += 1; continue
        if line.startswith("```"):
            if in_code:
                p=doc.add_paragraph("\n".join(code)); p.style="No Spacing"; code=[]; in_code=False
            else: in_code=True
            i += 1; continue
        if in_code: code.append(line); i += 1; continue
        if line.startswith("|") and i+1 < len(lines) and re.match(r"^\|?[ :\-|]+\|", lines[i+1]):
            rows=[]
            while i < len(lines) and lines[i].startswith("|"):
                cells=[c.strip() for c in lines[i].strip("|").split("|")]
                if not all(re.fullmatch(r":?-+:?", c.replace(" ","")) for c in cells): rows.append(cells)
                i += 1
            if rows:
                table=doc.add_table(rows=len(rows), cols=max(map(len,rows))); table.style="Table Grid"; table.autofit=False
                for rr,row in enumerate(rows):
                    for cc,val in enumerate(row): table.cell(rr,cc).text=re.sub(r"[*`]", "", val)
                column_width = Mm(116 / max(map(len, rows)))
                for row in table.rows:
                    for cell in row.cells: cell.width = column_width
                header_properties = table.rows[0]._tr.get_or_add_trPr()
                repeat_header = OxmlElement("w:tblHeader"); repeat_header.set(qn("w:val"), "true"); header_properties.append(repeat_header)
                for cell in table.rows[0].cells:
                    for run in cell.paragraphs[0].runs: run.bold=True
            continue
        m=re.match(r"^(#{1,3})\s+(.*)",line)
        if m: doc.add_heading(m.group(2), level=len(m.group(1))); i+=1; continue
        m=re.match(r"^\s*[-*]\s+(.*)",line)
        if m: add_inline(doc.add_paragraph(style="List Bullet"),m.group(1)); i+=1; continue
        m=re.match(r"^\s*\d+\.\s+(.*)",line)
        if m: add_inline(doc.add_paragraph(style="List Number"),m.group(1)); i+=1; continue
        if line.strip(): add_inline(doc.add_paragraph(), line)
        i += 1
    props=doc.core_properties; props.title=meta.get("title",BASENAME); props.subject=meta.get("subtitle","")
    doc.save(target)

def build_epub(meta: dict[str, str], body_html: str, target: Path, diagrams):
    from ebooklib import epub
    book=epub.EpubBook(); book.set_identifier("antifraud-ultimate-guide-2026-ru")
    book.set_title(meta.get("title",BASENAME)); book.set_language("ru")
    book.add_author(meta.get("author","")); css=epub.EpubItem(uid="style",file_name="style/book.css",media_type="text/css",content=CSS.encode())
    book.add_item(css); chapters=[]
    for d in diagrams:
        book.add_item(epub.EpubItem(uid=Path(d["svg"]).stem, file_name=f"assets/{Path(d['svg']).name}", media_type="image/svg+xml", content=Path(d["svg"]).read_bytes()))
        body_html=body_html.replace(Path(d["svg"]).as_posix(), f"../assets/{Path(d['svg']).name}")
    chunks=re.split(r"(?=<h1(?:\s|>))", body_html)
    for idx,chunk in enumerate(c for c in chunks if c.strip()):
        heading=re.search(r"<h1[^>]*>(.*?)</h1>",chunk,re.S)
        chapter_title=re.sub("<.*?>","",heading.group(1)) if heading else f"Раздел {idx+1}"
        ch=epub.EpubHtml(title=chapter_title,file_name=f"section_{idx+1:02}.xhtml",lang="ru")
        ch.content=chunk; ch.add_item(css); book.add_item(ch); chapters.append(ch)
    book.toc=tuple(chapters); book.spine=["nav",*chapters]
    book.add_item(epub.EpubNcx()); book.add_item(epub.EpubNav()); epub.write_epub(str(target),book,{})

def sha256(path: Path) -> str:
    h=hashlib.sha256(); h.update(path.read_bytes()); return h.hexdigest()

def verify(targets: list[Path]):
    import fitz
    pdf=fitz.open(targets[0]); assert len(pdf)>0 and all(p.rect.width>0 for p in pdf)
    pdf.close()
    for path in targets[1:]:
        with zipfile.ZipFile(path) as z: assert z.testzip() is None

def main():
    ap=argparse.ArgumentParser(); ap.add_argument("--clean",action="store_true"); args=ap.parse_args()
    require_modules()
    if args.clean and DIST.exists(): shutil.rmtree(DIST)
    DIST.mkdir(parents=True,exist_ok=True)
    meta,raw_body=strip_frontmatter(SOURCE.read_text(encoding="utf-8")); rendered_body,diagrams=render_mermaid(raw_body)
    body_html=markdown_html("[TOC]\n\n"+rendered_body); doc_html=full_html(meta,body_html)
    targets=[DIST/f"{BASENAME}.pdf",DIST/f"{BASENAME}.docx",DIST/f"{BASENAME}.epub"]
    build_pdf(doc_html,targets[0]); build_docx(meta,raw_body,targets[1],diagrams.copy()); build_epub(meta,body_html,targets[2],diagrams); verify(targets)
    manifest={"source":str(SOURCE.relative_to(ROOT)),"source_sha256":sha256(SOURCE),"artifacts":[{"file":p.name,"sha256":sha256(p),"bytes":p.stat().st_size} for p in targets]}
    (DIST/"manifest.json").write_text(json.dumps(manifest,ensure_ascii=False,indent=2)+"\n",encoding="utf-8")
    for p in targets: print(f"built {p.name}: {p.stat().st_size} bytes")
    print("verified PDF pages and ZIP CRC for DOCX/EPUB")
if __name__=="__main__": main()
