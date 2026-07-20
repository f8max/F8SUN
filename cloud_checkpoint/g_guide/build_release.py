#!/usr/bin/env python3
"""Build the reader edition in DOCX, PDF and EPUB from canonical UTF-8 sources."""
from __future__ import annotations

import html
import json
import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_BREAK
from docx.shared import Cm, Pt
from ebooklib import epub
import markdown
from weasyprint import HTML

ROOT = Path(__file__).parent
SRC = ROOT / "sources"
OUT = ROOT / "output"
TITLE = "СИМВОЛЫ: атлас видимого и невидимого"


def atlas_md() -> str:
    data = json.loads((SRC / "symbol_cards.json").read_text("utf-8"))
    blocks = []
    for card in data["карточки"]:
        blocks += [
            f"## {card['название']}",
            f"**Семейство:** {card['семейство']}  ",
            card["простое_ядро"],
            f"**Напряжение:** {card['напряжение_две_стороны']}  ",
            f"**Контексты:** {', '.join(card['контексты'])}.  ",
            f"**Авторский мост:** {card['авторская_интерпретация']}  ",
            f"**Граница чтения:** {card['типичная_ошибка_читателя']}",
        ]
    return "\n\n".join(blocks)


def glossary_md() -> str:
    rows = json.loads((SRC / "glossary.json").read_text("utf-8"))
    return "\n\n".join(
        f"## {row['term']}\n\n{row['definition']}\n\n**Не путать:** {row['not_same']}"
        for row in rows
    )


def zodiac_md() -> str:
    rows = json.loads((SRC / "zodiac_spreads.json").read_text("utf-8"))
    return "\n\n".join(
        f"## {row['name']}: {row['gesture']}\n\n{row['lead']}\n\n"
        f"**Видимое:** {row['visible']}  \n**Напряжение:** {row['tension']}  \n"
        f"**Нити:** {row['threads']}  \n**Граница:** {row['boundary']}  \n"
        f"**Вопрос:** {row['question']}"
        for row in rows
    )


def canonical_markdown() -> str:
    body = (SRC / "encyclopedia_manuscript.md").read_text("utf-8")
    body = (body.replace("{{SYMBOL_ATLAS}}", atlas_md())
                .replace("{{GLOSSARY}}", glossary_md())
                .replace("{{ZODIAC_ATLAS}}", zodiac_md()))
    visual_labels = {
        "02_alphabet_forms": "Алфавит форм: точка → линия → круг → пересечение",
        "03_polarity_axes": "Оси полярности: не выбор победителя, а поле напряжения",
        "04_cycle_of_meaning": "Цикл смысла: вещь → традиция → толкование → личный отклик",
        "05_zodiac_wheel": "Зодиакальное колесо: двенадцать жестов вокруг общего пояса",
        "06_reading_method": "Метод чтения: увидеть → назвать действие → проверить контекст → оставить остаток",
        "07_elemental_grammar": "Грамматика стихий: течь, гореть, веять, удерживать",
        "08_correspondence_orbit": "Орбита соответствий: мосты с разными основаниями не сливаются в тождество",
        "09_gemini_breakdown": "Близнецы: две вертикали, связь и пространство между",
        "10_bow_arrow_breakdown": "Лук и стрела: напряжение → отпускание → траектория → цель",
        "11_capricorn_breakdown": "Козерог: земная опора и водная глубина гибридной формы",
        "12_aquarius_breakdown": "Водолей: сосуд, поток и воздушный принцип передачи",
        "13_cancer_breakdown": "Рак: мягкое внутреннее, твёрдая граница, боковое движение",
        "14_north_star_breakdown": "Северная звезда: видимое вращение неба вокруг практического ориентира",
    }
    body = re.sub(
        r"\{\{VISUAL:([^}]+)\}\}",
        lambda m: f"> **Схема для взгляда.** {visual_labels.get(m.group(1), m.group(1))}.",
        body,
    )
    nav = (SRC / "reader_navigation_and_index.md").read_text("utf-8")
    return f"# {TITLE}\n\n{nav}\n\n{body}"


def html_document(md: str) -> str:
    content = markdown.markdown(md, extensions=["extra", "sane_lists"])
    content = re.sub(r"<p>:::callout[^>]*</p>", '<aside class="callout">', content)
    content = content.replace("<p>:::</p>", "</aside>")
    return f"""<!doctype html><html lang="ru"><head><meta charset="utf-8">
<title>{html.escape(TITLE)}</title><style>
@page {{ size: A5; margin: 18mm 16mm 20mm; @bottom-center {{content: counter(page)}} }}
body {{ font-family: 'DejaVu Serif', serif; font-size: 10.2pt; line-height: 1.48; color:#20252b }}
h1 {{ page-break-before: always; color:#183e4b; font-size:24pt; line-height:1.08 }}
h1:first-child {{ page-break-before: avoid; margin-top:45mm; font-size:30pt }}
h2 {{ color:#9b4f36; font-size:16pt; margin-top:1.4em; page-break-after:avoid }}
h3 {{ color:#315965; font-size:12pt; page-break-after:avoid }}
p {{ orphans:3; widows:3; text-align:justify; hyphens:auto }}
blockquote,.callout {{ border-left:3px solid #c39255; padding:.3em 1em; background:#f4efe5 }}
table {{ border-collapse:collapse; font-size:8.5pt }} td,th {{ border:1px solid #aaa; padding:4px }}
a {{ color:#315965; text-decoration:none }}
</style></head><body>{content}</body></html>"""


def build_docx(md: str, path: Path) -> None:
    doc = Document()
    sec = doc.sections[0]
    sec.page_height, sec.page_width = Cm(21), Cm(14.8)
    sec.left_margin = sec.right_margin = Cm(1.7)
    normal = doc.styles["Normal"]
    normal.font.name, normal.font.size = "DejaVu Serif", Pt(10)
    for line in md.splitlines():
        line = line.strip()
        if not line:
            continue
        if line.startswith("# "):
            if len(doc.paragraphs) > 1:
                doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
            doc.add_heading(line[2:], 0)
        elif line.startswith("## "):
            doc.add_heading(line[3:], 1)
        elif line.startswith("### "):
            doc.add_heading(line[4:], 2)
        elif line.startswith(("- ", "* ")):
            doc.add_paragraph(re.sub(r"\*\*(.*?)\*\*", r"\1", line[2:]), style="List Bullet")
        elif re.match(r"\d+\. ", line):
            doc.add_paragraph(re.sub(r"^\d+\. ", "", line), style="List Number")
        elif not line.startswith(":::"):
            doc.add_paragraph(re.sub(r"\*\*(.*?)\*\*", r"\1", line))
    doc.core_properties.title = TITLE
    doc.core_properties.language = "ru-RU"
    doc.save(path)


def build_epub(md: str, path: Path) -> None:
    book = epub.EpubBook()
    book.set_identifier("g-guide-symbols-2026")
    book.set_title(TITLE)
    book.set_language("ru")
    css = epub.EpubItem(uid="style", file_name="style.css", media_type="text/css",
                        content="body{font-family:serif;line-height:1.5}h1,h2{color:#315965}aside{border-left:3px solid #c39255;padding:1em}")
    book.add_item(css)
    chunks = re.split(r"(?=^# )", md, flags=re.M)
    chapters = []
    for i, chunk in enumerate(filter(str.strip, chunks)):
        title_match = re.search(r"^# (.+)", chunk)
        title = title_match.group(1) if title_match else TITLE
        chapter = epub.EpubHtml(title=title, file_name=f"part_{i:02d}.xhtml", lang="ru")
        chapter.content = markdown.markdown(chunk, extensions=["extra"])
        chapter.add_item(css)
        book.add_item(chapter)
        chapters.append(chapter)
    book.toc = tuple(chapters)
    book.spine = ["nav", *chapters]
    book.add_item(epub.EpubNcx())
    book.add_item(epub.EpubNav())
    epub.write_epub(path, book)


def main() -> None:
    OUT.mkdir(exist_ok=True)
    md = canonical_markdown()
    (OUT / "G_GUIDE_SYMBOLS_RELEASE.md").write_text(md, "utf-8")
    html_text = html_document(md)
    HTML(string=html_text, base_url=str(ROOT)).write_pdf(OUT / "G_GUIDE_SYMBOLS_RELEASE.pdf")
    build_docx(md, OUT / "G_GUIDE_SYMBOLS_RELEASE.docx")
    build_epub(md, OUT / "G_GUIDE_SYMBOLS_RELEASE.epub")
    print(f"Built {len(md):,} characters in {OUT}")


if __name__ == "__main__":
    main()
